import copy
from docx import Document
from docx.oxml.ns import qn
from lxml import etree as ET

meetings = [
    (
        11, '21 April, 2026', '10:30 AM',
        [
            'Set up Firebase project and connect the Flutter app with Firestore and Authentication modules.',
            'Define the project folder structure and finalize all package dependencies in pubspec.yaml.',
            'Plan the development roadmap covering the Flutter app, firmware, and ML components.',
        ]
    ),
    (
        12, '28 April, 2026', '11:00 AM',
        [
            'Implement bottom navigation bar and screen routing for all five main app sections.',
            'Build scaffold layouts for Home, Map, Alerts, Profile, and Settings screens.',
            'Review screen wireframes against the Velvet Night design system and adjust spacing.',
        ]
    ),
    (
        13, '07 May, 2026', '10:00 AM',
        [
            'Implement Firebase Authentication for email/password login and signup flows.',
            'Add form validation, error handling, and session persistence on app relaunch.',
            'Test login, signup, and logout functionality on emulator and physical device.',
        ]
    ),
    (
        14, '13 May, 2026', '11:30 AM',
        [
            'Integrate Google Maps Flutter plugin and display live child location on the map screen.',
            'Set up a Firestore real-time listener to stream device GPS coordinates to the app.',
            'Add custom map markers and a child avatar overlay for the live tracking view.',
        ]
    ),
    (
        15, '20 May, 2026', '10:30 AM',
        [
            'Implement safe zone creation with a draggable circle overlay on the map.',
            'Store safe zone boundaries (center coordinates and radius) in Firestore per user.',
            'Trigger a local notification when the child location exits the defined safe zone.',
        ]
    ),
    (
        16, '26 May, 2026', '11:00 AM',
        [
            'Write ESP32 firmware to read GPS coordinates from the Neo-6M module via UART.',
            'Implement BLE advertisement to broadcast device ID and GPS data to the paired phone.',
            'Test firmware on the hardware prototype and verify BLE data reception in Flutter.',
        ]
    ),
    (
        17, '03 June, 2026', '10:30 AM',
        [
            'Integrate TensorFlow Lite model on ESP32 for activity detection using MPU-6050 data.',
            'Implement a preprocessing pipeline to normalize accelerometer readings before inference.',
            'Validate model output labels (idle, walking, running, fall) against real movement samples.',
        ]
    ),
    (
        18, '09 June, 2026', '11:00 AM',
        [
            'Set up Firebase Cloud Messaging for push notifications on safety alert events.',
            'Trigger alerts for safe zone breach, fall detection, and low battery conditions.',
            'Build an in-app alert history screen displaying timestamp and alert type per event.',
        ]
    ),
    (
        19, '16 June, 2026', '10:00 AM',
        [
            'Conduct end-to-end testing of the ESP32-to-Flutter BLE communication pipeline.',
            'Test Firestore real-time sync for location updates and alert triggering under load.',
            'Document all bugs found during integration testing and assign fixes for the next session.',
        ]
    ),
    (
        20, '18 June, 2026', '11:30 AM',
        [
            'Review complete system functionality covering hardware, firmware, and mobile app.',
            'Prepare demo walkthrough script and finalize project documentation for FYP submission.',
            'Discuss remaining issues and confirm overall readiness for the final evaluation.',
        ]
    ),
]


def add_bullet_paragraph(tc, text, template_para):
    new_p = copy.deepcopy(template_para._element)
    for r in new_p.findall(qn('w:r')):
        new_p.remove(r)
    for ins in new_p.findall(qn('w:ins')):
        new_p.remove(ins)
    r_elem = copy.deepcopy(template_para.runs[0]._element)
    for t in r_elem.findall(qn('w:t')):
        r_elem.remove(t)
    t_elem = ET.SubElement(r_elem, qn('w:t'))
    t_elem.text = text
    t_elem.set('{http://www.w3.org/XML/1998/namespace}space', 'preserve')
    new_p.append(r_elem)
    tc.append(new_p)


for meeting_no, date, time_str, action_items in meetings:
    doc = Document('forms/Meeting Form 10.docx')
    table = doc.tables[0]

    # 1. Meeting Number (row 6, col 1)
    cell = table.rows[6].cells[1]
    for p in cell.paragraphs:
        for run in p.runs:
            run.text = ''
    if cell.paragraphs[0].runs:
        cell.paragraphs[0].runs[0].text = str(meeting_no)
    else:
        cell.paragraphs[0].add_run(str(meeting_no))

    # 2. Date (row 6, col 3)
    cell_date = table.rows[6].cells[3]
    for p in cell_date.paragraphs:
        for run in p.runs:
            run.text = ''
    if cell_date.paragraphs[0].runs:
        cell_date.paragraphs[0].runs[0].text = date
    else:
        cell_date.paragraphs[0].add_run(date)

    # 3. Time (row 6, col 6)
    cell_time = table.rows[6].cells[6]
    for p in cell_time.paragraphs:
        for run in p.runs:
            run.text = ''
    if cell_time.paragraphs[0].runs:
        cell_time.paragraphs[0].runs[0].text = time_str
    else:
        cell_time.paragraphs[0].add_run(time_str)

    # 4. Action Items (row 9) - replace with 3 bullet points
    action_cell = table.rows[9].cells[0]
    template_para = action_cell.paragraphs[0]
    tc = action_cell._tc
    for p in tc.findall(qn('w:p')):
        tc.remove(p)
    for item_text in action_items:
        add_bullet_paragraph(tc, item_text, template_para)
    # trailing empty paragraph
    end_p = copy.deepcopy(template_para._element)
    for child in list(end_p):
        if child.tag != qn('w:pPr'):
            end_p.remove(child)
    pPr = end_p.find(qn('w:pPr'))
    if pPr is not None:
        numPr = pPr.find(qn('w:numPr'))
        if numPr is not None:
            pPr.remove(numPr)
        pStyle = pPr.find(qn('w:pStyle'))
        if pStyle is not None:
            pStyle.set(qn('w:val'), 'NoSpacing')
    tc.append(end_p)

    # 5. Progress section (row 12) - leave blank (already empty in template)

    filename = f'forms/Meeting Form {meeting_no}.docx'
    doc.save(filename)
    print(f'Saved: {filename}')

print('All done!')
