from docx import Document
from docx.shared import Pt, RGBColor, Cm, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

# All 10 meetings: (number, date, previous_items, progress_notes)
# previous_items = action items from the PREVIOUS meeting
# progress_notes = what was accomplished on those items

data = [
    {
        'form': 11,
        'date': '21 April, 2026',
        'previous_items': [
            'Review and refine block, use-case, and sequence diagrams.',
            'Discuss whether diagrams clearly represent system functionality.',
            'Align diagrams with the defended concept and panel feedback.',
            'Discuss how these diagrams guide the implementation phase and finalize design-phase material for submission.',
        ],
        'progress': [
            'All diagrams reviewed, refined, and updated to reflect final system architecture; sequence diagram revised to include BLE communication flow.',
            'Diagrams confirmed to clearly represent system functionality; minor corrections applied to use-case diagram following supervisor feedback.',
            'Design-phase documentation finalized and organized for FYP records and submission.',
        ],
    },
    {
        'form': 12,
        'date': '28 April, 2026',
        'previous_items': [
            'Set up Firebase project and connect the Flutter app with Firestore and Authentication modules.',
            'Define the project folder structure and finalize all package dependencies in pubspec.yaml.',
            'Plan the development roadmap covering the Flutter app, firmware, and ML components.',
        ],
        'progress': [
            'Firebase project created and successfully connected to the Flutter app; Firestore and Authentication initialized with proper security rules.',
            'Feature-based project folder structure defined; pubspec.yaml updated with all required dependencies and version constraints.',
            'Development roadmap drafted and shared with supervisor, covering all three phases: Flutter app, ESP32 firmware, and TFLite ML model.',
        ],
    },
    {
        'form': 13,
        'date': '07 May, 2026',
        'previous_items': [
            'Implement bottom navigation bar and screen routing for all five main app sections.',
            'Build scaffold layouts for Home, Map, Alerts, Profile, and Settings screens.',
            'Review screen wireframes against the Velvet Night design system and adjust spacing.',
        ],
        'progress': [
            'Bottom navigation bar implemented with persistent routing across all five app sections using Go Router.',
            'Scaffold layouts built for all screens with consistent theming and placeholder widgets.',
            'Wireframes reviewed against design tokens; spacing, font sizes, and colour values adjusted to match the Velvet Night design system.',
        ],
    },
    {
        'form': 14,
        'date': '13 May, 2026',
        'previous_items': [
            'Implement Firebase Authentication for email/password login and signup flows.',
            'Add form validation, error handling, and session persistence on app relaunch.',
            'Test login, signup, and logout functionality on emulator and physical device.',
        ],
        'progress': [
            'Firebase Authentication integrated with fully working email/password login and signup screens.',
            'Form validation, descriptive error messages, and session persistence implemented and working correctly on relaunch.',
            'Login, signup, and logout flows tested and passing on both Android emulator and a physical device.',
        ],
    },
    {
        'form': 15,
        'date': '20 May, 2026',
        'previous_items': [
            'Integrate Google Maps Flutter plugin and display live child location on the map screen.',
            'Set up a Firestore real-time listener to stream device GPS coordinates to the app.',
            'Add custom map markers and a child avatar overlay for the live tracking view.',
        ],
        'progress': [
            'Google Maps Flutter plugin integrated; live child location displayed and updating on the map screen.',
            'Firestore real-time listener configured and streaming GPS coordinates with low latency.',
            'Custom map markers and child avatar overlay implemented and rendering correctly on the tracking view.',
        ],
    },
    {
        'form': 16,
        'date': '26 May, 2026',
        'previous_items': [
            'Implement safe zone creation with a draggable circle overlay on the map.',
            'Store safe zone boundaries (center coordinates and radius) in Firestore per user.',
            'Trigger a local notification when the child location exits the defined safe zone.',
        ],
        'progress': [
            'Safe zone creation implemented with a fully draggable and resizable circle overlay on the map.',
            'Safe zone boundaries stored in Firestore under the user document and correctly retrieved on app launch.',
            'Local notification triggered and tested; alert fires reliably when the child location moves outside the defined boundary.',
        ],
    },
    {
        'form': 17,
        'date': '03 June, 2026',
        'previous_items': [
            'Write ESP32 firmware to read GPS coordinates from the Neo-6M module via UART.',
            'Implement BLE advertisement to broadcast device ID and GPS data to the paired phone.',
            'Test firmware on the hardware prototype and verify BLE data reception in Flutter.',
        ],
        'progress': [
            'ESP32 firmware written and successfully reading GPS coordinates from the Neo-6M module with NMEA sentence parsing.',
            'BLE advertisement implemented; device ID and GPS payload broadcasting confirmed using a BLE scanner app.',
            'Firmware tested on the hardware prototype; Flutter app receiving and correctly parsing the BLE data packets.',
        ],
    },
    {
        'form': 18,
        'date': '09 June, 2026',
        'previous_items': [
            'Integrate TensorFlow Lite model on ESP32 for activity detection using MPU-6050 data.',
            'Implement a preprocessing pipeline to normalize accelerometer readings before inference.',
            'Validate model output labels (idle, walking, running, fall) against real movement samples.',
        ],
        'progress': [
            'TFLite model successfully integrated on ESP32 and running inference on live MPU-6050 accelerometer data.',
            'Preprocessing pipeline implemented; raw accelerometer values normalized and windowed before being passed to the model.',
            'Model output validated against real-world movement samples; fall detection and activity classification performing within acceptable accuracy.',
        ],
    },
    {
        'form': 19,
        'date': '16 June, 2026',
        'previous_items': [
            'Set up Firebase Cloud Messaging for push notifications on safety alert events.',
            'Trigger alerts for safe zone breach, fall detection, and low battery conditions.',
            'Build an in-app alert history screen displaying timestamp and alert type per event.',
        ],
        'progress': [
            'Firebase Cloud Messaging configured and push notifications working end-to-end for all safety alert types.',
            'All three alert conditions (safe zone breach, fall detection, low battery) triggering correctly in test scenarios.',
            'Alert history screen built and displaying entries with timestamp and alert type, with data persisted in Firestore.',
        ],
    },
    {
        'form': 20,
        'date': '18 June, 2026',
        'previous_items': [
            'Conduct end-to-end testing of the ESP32-to-Flutter BLE communication pipeline.',
            'Test Firestore real-time sync for location updates and alert triggering under load.',
            'Document all bugs found during integration testing and assign fixes for the next session.',
        ],
        'progress': [
            'End-to-end BLE communication pipeline tested; a packet-loss issue identified and resolved in the BLE handler.',
            'Firestore real-time sync tested under load; location updates and alert delivery performing within acceptable thresholds.',
            'All integration bugs documented; fixes applied and system confirmed stable ahead of final evaluation.',
        ],
    },
]


def set_cell_bg(cell, hex_color):
    tc = cell._tc
    tcPr = tc.find(qn('w:tcPr'))
    if tcPr is None:
        tcPr = OxmlElement('w:tcPr')
        tc.insert(0, tcPr)
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), hex_color)
    tcPr.append(shd)


def set_cell_borders(cell, top=None, bottom=None, left=None, right=None):
    tc = cell._tc
    tcPr = tc.find(qn('w:tcPr'))
    if tcPr is None:
        tcPr = OxmlElement('w:tcPr')
        tc.insert(0, tcPr)
    tcBorders = OxmlElement('w:tcBorders')
    for side, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        if val is not None:
            border = OxmlElement(f'w:{side}')
            border.set(qn('w:val'), val.get('val', 'single'))
            border.set(qn('w:sz'), str(val.get('sz', 4)))
            border.set(qn('w:space'), '0')
            border.set(qn('w:color'), val.get('color', '000000'))
            tcBorders.append(border)
    tcPr.append(tcBorders)


doc = Document()

# Page margins
section = doc.sections[0]
section.top_margin = Cm(1.5)
section.bottom_margin = Cm(1.5)
section.left_margin = Cm(2.0)
section.right_margin = Cm(2.0)

# Title paragraph
title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = title.add_run('AEGIS – Child Safety Band')
run.bold = True
run.font.size = Pt(14)

sub = doc.add_paragraph()
sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
run2 = sub.add_run('Progress on Previous Action Items — Meeting Forms 11 to 20')
run2.bold = True
run2.font.size = Pt(11)

doc.add_paragraph()

for entry in data:
    form_no = entry['form']
    date = entry['date']
    prev_items = entry['previous_items']
    progress = entry['progress']

    # ── Meeting header table (1 row, 2 cols) ──────────────────────────────
    hdr_table = doc.add_table(rows=1, cols=2)
    hdr_table.style = 'Table Grid'
    hdr_table.alignment = WD_TABLE_ALIGNMENT.CENTER

    c0 = hdr_table.rows[0].cells[0]
    c1 = hdr_table.rows[0].cells[1]

    set_cell_bg(c0, '1F3864')  # dark navy
    set_cell_bg(c1, '1F3864')

    p0 = c0.paragraphs[0]
    p0.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r0 = p0.add_run(f'Meeting Form No.  {form_no}')
    r0.bold = True
    r0.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    r0.font.size = Pt(10)

    p1 = c1.paragraphs[0]
    p1.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r1 = p1.add_run(f'Date:  {date}')
    r1.bold = True
    r1.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    r1.font.size = Pt(10)

    # Set column widths
    hdr_table.columns[0].width = Cm(9)
    hdr_table.columns[1].width = Cm(8)

    # ── Action Items from Previous Meeting ────────────────────────────────
    ai_table = doc.add_table(rows=1 + len(prev_items), cols=1)
    ai_table.style = 'Table Grid'
    ai_table.alignment = WD_TABLE_ALIGNMENT.CENTER

    # Sub-header row
    sh_cell = ai_table.rows[0].cells[0]
    set_cell_bg(sh_cell, 'D6E4F0')
    sh_p = sh_cell.paragraphs[0]
    sh_r = sh_p.add_run('Action Items from Previous Meeting')
    sh_r.bold = True
    sh_r.font.size = Pt(9.5)

    for i, item in enumerate(prev_items):
        row_cell = ai_table.rows[i + 1].cells[0]
        row_p = row_cell.paragraphs[0]
        row_p.paragraph_format.left_indent = Pt(6)
        row_r = row_p.add_run(f'•  {item}')
        row_r.font.size = Pt(9)

    # ── Progress on Previous Items ────────────────────────────────────────
    prog_table = doc.add_table(rows=1 + len(progress), cols=1)
    prog_table.style = 'Table Grid'
    prog_table.alignment = WD_TABLE_ALIGNMENT.CENTER

    # Sub-header row
    ph_cell = prog_table.rows[0].cells[0]
    set_cell_bg(ph_cell, 'D5F5E3')
    ph_p = ph_cell.paragraphs[0]
    ph_r = ph_p.add_run('Progress on Previous Items')
    ph_r.bold = True
    ph_r.font.size = Pt(9.5)

    for i, note in enumerate(progress):
        row_cell = prog_table.rows[i + 1].cells[0]
        row_p = row_cell.paragraphs[0]
        row_p.paragraph_format.left_indent = Pt(6)
        row_r = row_p.add_run(f'•  {note}')
        row_r.font.size = Pt(9)

    # ── Supervisor sign-off ───────────────────────────────────────────────
    sig_table = doc.add_table(rows=1, cols=2)
    sig_table.style = 'Table Grid'
    sig_table.alignment = WD_TABLE_ALIGNMENT.CENTER

    sc0 = sig_table.rows[0].cells[0]
    sc1 = sig_table.rows[0].cells[1]

    sp0 = sc0.paragraphs[0]
    sp0.add_run("Supervisor's Comments:  ").font.size = Pt(9)
    sc0.add_paragraph().add_run('').font.size = Pt(9)  # blank line

    sp1 = sc1.paragraphs[0]
    sp1.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    sp1.add_run("Supervisor's Signature:  ___________________").font.size = Pt(9)

    sig_table.columns[0].width = Cm(10)
    sig_table.columns[1].width = Cm(7)

    # Spacer between sections
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(4)

out_path = 'forms/Progress Notes (Forms 11-20).docx'
doc.save(out_path)
print(f'Saved: {out_path}')
