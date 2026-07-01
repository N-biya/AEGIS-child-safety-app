"""Generate the AEGIS firmware -> Firebase integration guide as a .docx."""
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

OUT = r"E:\FypApp\AEGIS_Firmware_Firebase_Integration.docx"

ACCENT = RGBColor(0x2B, 0x2D, 0x6E)   # velvet night-ish
MONO_BG = RGBColor(0xF2, 0xF2, 0xF6)


def code_block(doc, text):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.name = "Consolas"
    run.font.size = Pt(9.5)
    p.paragraph_format.left_indent = Inches(0.2)
    p.paragraph_format.space_after = Pt(6)
    return p


def kv_table(doc, rows, headers=("Field", "Value")):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Light Grid Accent 1"
    for i, h in enumerate(headers):
        c = t.rows[0].cells[i]
        c.text = h
        for r in c.paragraphs[0].runs:
            r.font.bold = True
    for row in rows:
        cells = t.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = str(val)
            if i == len(headers) - 1:
                for r in cells[i].paragraphs[0].runs:
                    r.font.name = "Consolas"
                    r.font.size = Pt(9.5)
    doc.add_paragraph()
    return t


def h(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    for r in p.runs:
        r.font.color.rgb = ACCENT
    return p


def main():
    doc = Document()

    # Title
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = title.add_run("AEGIS Wristband")
    r.font.size = Pt(26); r.font.bold = True; r.font.color.rgb = ACCENT
    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rs = sub.add_run("Firmware → Firebase Integration Guide")
    rs.font.size = Pt(14); rs.font.color.rgb = RGBColor(0x66, 0x66, 0x66)
    doc.add_paragraph()

    # 1. Overview
    h(doc, "1. How data flows")
    doc.add_paragraph(
        "The wristband sends vitals straight to the cloud over Wi-Fi — there is no "
        "intermediate phone or server. The ESP32 connects to Wi-Fi, authenticates to "
        "Firebase using a dedicated device account, and writes one document per reading "
        "directly into Cloud Firestore over HTTPS (TLS). The parent app reads those same "
        "documents in real time."
    )
    code_block(doc,
        "ESP32  --Wi-Fi/HTTPS-->  Firebase (Cloud Firestore)  -->  Parent App")
    doc.add_paragraph(
        "If Wi-Fi drops, readings are encrypted and held in an on-device offline queue, "
        "then flushed automatically when the connection returns — no data is lost."
    )

    # 2. Credentials
    h(doc, "2. Configuration values (secrets.h)")
    doc.add_paragraph(
        "Copy secrets.example.h to secrets.h and fill in the values below. "
        "secrets.h must never be committed to version control (it is already git-ignored)."
    )
    kv_table(doc, [
        ("FIREBASE_PROJECT_ID",    "aegis-child-safety-app"),
        ("FIREBASE_API_KEY",       "AIzaSyC-_Frb8dG9Oi4G6aKZEIb5teFbaZDHFVA"),
        ("FIREBASE_HOST",          "firestore.googleapis.com   (leave unchanged)"),
        ("FIREBASE_USER_EMAIL",    "band-wajeehaa@aegis.device"),
        ("FIREBASE_USER_PASSWORD", "za8jwpxUV7hmTakS86kOKw"),
        ("CHILD_ID",               "ecc086c4-215c-4661-b6c4-693cadb54d16"),
        ("USER_ID",                "(not used in paths — may be left blank)"),
        ("WIFI_SSID",              "<the Wi-Fi network the band will use>"),
        ("WIFI_PASSWORD",          "<that network's password>"),
    ])
    doc.add_paragraph(
        "The device account and its permissions are already set up in the project; the "
        "email and password above are ready to use as-is. Only the two Wi-Fi values need "
        "to be supplied to match the deployment location."
    )

    # 3. Where data lands
    h(doc, "3. Where the data is written")
    doc.add_paragraph("Vitals are written once per upload interval to:")
    code_block(doc, "children/{CHILD_ID}/vitals/{autoId}")
    doc.add_paragraph("Alerts (stress, geofence breach, low SpO2, etc.) are written to:")
    code_block(doc, "children/{CHILD_ID}/alerts/{autoId}")

    # 4. Vitals schema
    h(doc, "4. Vitals document fields", level=2)
    kv_table(doc, [
        ("hr",        "number  — heart rate (BPM)"),
        ("spo2",      "number  — blood oxygen (%)"),
        ("gsr",       "number  — skin conductance (µS)"),
        ("temp",      "number  — skin temperature (°C)"),
        ("movement",  "number  — acceleration magnitude"),
        ("status",    'string  — "NORMAL" | "ELEVATED" | "STRESS"'),
        ("lat",       "number  — GPS latitude (0 when no fix)"),
        ("lng",       "number  — GPS longitude (0 when no fix)"),
        ("timestamp", 'string  — ISO-8601 UTC, e.g. "2026-06-29T09:31:43Z"'),
        ("battery",   "number  — battery percentage (optional)"),
    ], headers=("Field", "Type / meaning"))

    # 5. Alert schema
    h(doc, "5. Alert document fields", level=2)
    kv_table(doc, [
        ("type",      'string  — "STRESS" | "GEOFENCE" | "SPO2" | ...'),
        ("timestamp", "string  — ISO-8601 UTC"),
        ("vitals",    "map     — { hr: number, spo2: number }"),
        ("location",  "map     — { lat: number, lng: number }"),
        ("resolved",  "boolean — false on creation"),
    ], headers=("Field", "Type / meaning"))

    # 6. Setup steps
    h(doc, "6. Setup checklist")
    for i, step in enumerate([
        "Copy secrets.example.h to secrets.h and paste in the values from section 2.",
        "Set WIFI_SSID and WIFI_PASSWORD to the network the band will run on.",
        "Build and flash the firmware (the field/path changes only take effect after reflashing).",
        "Power on the band within Wi-Fi range. After the first upload, documents appear under "
        "children/{CHILD_ID}/vitals and the app dashboard updates live.",
    ], start=1):
        doc.add_paragraph(f"{i}. {step}")

    # 7. Notes
    h(doc, "7. Notes")
    for note in [
        "Time sync: the firmware pulls the time from an NTP server on connect so each "
        "document gets a real UTC timestamp the app can read. Allow a few seconds after "
        "Wi-Fi connects before the first upload.",
        "No endpoint URL is required — the Firebase client library handles the REST call "
        "to Firestore internally.",
        "All physiological and location data travels over HTTPS only; nothing sensitive is "
        "printed to the serial console in release builds.",
        "Keep secrets.h out of git. It contains the device password.",
    ]:
        p = doc.add_paragraph(style="List Bullet")
        p.add_run(note)

    doc.save(OUT)
    print("Saved:", OUT)


if __name__ == "__main__":
    main()
