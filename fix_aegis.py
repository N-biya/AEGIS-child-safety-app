#!/usr/bin/env python3
"""
fix_aegis.py
Input:  AEGIS_with_diagramss.docx  (original, untouched)
Output: AEGIS_fixed.docx

Fix 1: Remove the duplicate numPr from the Heading3 STYLE definition in styles.xml.
       This leaves only the paragraph-level numPr (numId=11, ilvl=0/1/2) so each
       heading renders a single number.

Fix 2: (a) Add SEQ Table caption paragraphs after each content table.
       (b) Keep the existing LOF TOC field intact.
       (c) Insert a List of Tables heading + { TOC \h \z \c "Table" } field.
"""
import sys, io, shutil, re, zipfile
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')

from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import lxml.etree as etree

INPUT  = r'e:\FypApp\AEGIS_with_diagramss.docx'
OUTPUT = r'e:\FypApp\AEGIS_fixed.docx'

# Table captions for doc.tables indices 1-7 (index 0 is empty 1×1)
TABLE_CAPTIONS = {
    1: 'Board of Examiners',
    2: 'Comparison of Existing Child Safety Solutions',
    3: 'Use Case Flow: View Real-Time Vitals',
    4: 'Use Case Flow: Monitor Heart Rate and SpO₂',
    5: 'Use Case Flow: Monitor Body Temperature',
    6: 'Use Case Flow: Monitor Stress Level (GSR)',
    7: 'Use Case Flow: View Live Location and History',
}
TNR = 'Times New Roman'

# ── XML helpers ───────────────────────────────────────────────────────────────

def make_run(text):
    """Simple run with Times New Roman 12 pt."""
    r = OxmlElement('w:r')
    rPr = OxmlElement('w:rPr')
    rf = OxmlElement('w:rFonts')
    rf.set(qn('w:ascii'), TNR); rf.set(qn('w:hAnsi'), TNR); rf.set(qn('w:cs'), TNR)
    rPr.append(rf)
    sz = OxmlElement('w:sz'); sz.set(qn('w:val'), '24'); rPr.append(sz)
    szc = OxmlElement('w:szCs'); szc.set(qn('w:val'), '24'); rPr.append(szc)
    r.append(rPr)
    t = OxmlElement('w:t'); t.set(qn('xml:space'), 'preserve'); t.text = text
    r.append(t)
    return r

def make_seq_field(label, cached_num):
    """
    Build:  begin → instrText " SEQ <label> \* ARABIC " → separate → cached → end
    Returns a list of w:r elements to append to a paragraph element.
    """
    W = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'
    runs = []

    r_begin = OxmlElement('w:r')
    fc = OxmlElement('w:fldChar'); fc.set(qn('w:fldCharType'), 'begin')
    r_begin.append(fc); runs.append(r_begin)

    r_instr = OxmlElement('w:r')
    it = OxmlElement('w:instrText')
    it.set(qn('xml:space'), 'preserve')
    it.text = f' SEQ {label} \\* ARABIC '
    r_instr.append(it); runs.append(r_instr)

    r_sep = OxmlElement('w:r')
    fc2 = OxmlElement('w:fldChar'); fc2.set(qn('w:fldCharType'), 'separate')
    r_sep.append(fc2); runs.append(r_sep)

    r_cache = OxmlElement('w:r')
    t = OxmlElement('w:t'); t.text = str(cached_num); r_cache.append(t)
    runs.append(r_cache)

    r_end = OxmlElement('w:r')
    fc3 = OxmlElement('w:fldChar'); fc3.set(qn('w:fldCharType'), 'end')
    r_end.append(fc3); runs.append(r_end)

    return runs

def make_toc_field(label):
    """
    Build a single paragraph element containing:
        { TOC \h \z \c "Label" }
    with style TableofFigures (matching the existing LOF paragraph's style).
    """
    p = OxmlElement('w:p')

    # pPr  — use TableofFigures style to match the existing LOF
    pPr = OxmlElement('w:pPr')
    ps = OxmlElement('w:pStyle'); ps.set(qn('w:val'), 'TableofFigures'); pPr.append(ps)
    p.append(pPr)

    # Complex field: begin → instrText → separate → end  (no cached result yet)
    r_begin = OxmlElement('w:r')
    fc_b = OxmlElement('w:fldChar'); fc_b.set(qn('w:fldCharType'), 'begin')
    r_begin.append(fc_b); p.append(r_begin)

    r_instr = OxmlElement('w:r')
    it = OxmlElement('w:instrText')
    it.set(qn('xml:space'), 'preserve')
    it.text = f' TOC \\h \\z \\c "{label}" '
    r_instr.append(it); p.append(r_instr)

    r_sep = OxmlElement('w:r')
    fc_s = OxmlElement('w:fldChar'); fc_s.set(qn('w:fldCharType'), 'separate')
    r_sep.append(fc_s); p.append(r_sep)

    # Placeholder text (will be replaced by Word on field update)
    r_ph = OxmlElement('w:r')
    t_ph = OxmlElement('w:t')
    t_ph.text = f'[Update field to see List of {label}s]'
    r_ph.append(t_ph); p.append(r_ph)

    r_end = OxmlElement('w:r')
    fc_e = OxmlElement('w:fldChar'); fc_e.set(qn('w:fldCharType'), 'end')
    r_end.append(fc_e); p.append(r_end)

    return p

def make_caption_para(label, seq_num, caption_text, bookmark_id, bookmark_name):
    """
    Build a Caption-style paragraph:
        <label> <SEQ field> <space><caption_text>
    with a bookmark wrapping the whole paragraph (so the TOC \h link jumps here).
    """
    p = OxmlElement('w:p')

    # pPr
    pPr = OxmlElement('w:pPr')
    ps = OxmlElement('w:pStyle'); ps.set(qn('w:val'), 'Caption'); pPr.append(ps)
    jc = OxmlElement('w:jc'); jc.set(qn('w:val'), 'center'); pPr.append(jc)
    p.append(pPr)

    # Bookmark start
    bkStart = OxmlElement('w:bookmarkStart')
    bkStart.set(qn('w:id'), str(bookmark_id))
    bkStart.set(qn('w:name'), bookmark_name)
    p.append(bkStart)

    # "Table " text run
    p.append(make_run(f'{label} '))

    # SEQ field runs
    for r in make_seq_field(label, seq_num):
        p.append(r)

    # space + caption title
    p.append(make_run(f' {caption_text}'))

    # Bookmark end
    bkEnd = OxmlElement('w:bookmarkEnd')
    bkEnd.set(qn('w:id'), str(bookmark_id))
    p.append(bkEnd)

    return p

def make_lot_heading():
    """Heading 2 paragraph: 'List of Tables'"""
    p = OxmlElement('w:p')
    pPr = OxmlElement('w:pPr')
    ps = OxmlElement('w:pStyle'); ps.set(qn('w:val'), 'Heading2'); pPr.append(ps)
    p.append(pPr)
    r = OxmlElement('w:r')
    t = OxmlElement('w:t'); t.text = 'List of Tables'; r.append(t)
    p.append(r)
    return p

# ── FIX 1: Remove numPr from Heading3 style definition ───────────────────────

def fix_heading3_numpr(doc):
    """
    Remove the numPr element from the Heading3 paragraph STYLE.
    This kills the style-layer numbering; each paragraph's own numPr(numId=11)
    will continue to provide the correct hierarchical numbers.
    Returns True if found and removed.
    """
    styles_elem = doc.part.styles._element
    for style in styles_elem.findall(qn('w:style')):
        sId = style.get(qn('w:styleId'))
        if sId == 'Heading3':
            pPr = style.find(qn('w:pPr'))
            if pPr is not None:
                numPr = pPr.find(qn('w:numPr'))
                if numPr is not None:
                    pPr.remove(numPr)
                    print('[OK] Removed numPr from Heading3 style definition')
                    return True
                else:
                    print('[INFO] Heading3 style has no numPr — nothing to remove')
            return False
    print('[WARN] Heading3 style not found')
    return False

# ── FIX 2: Add table captions + LOT ──────────────────────────────────────────

def get_max_bookmark_id(doc):
    """Find the highest bookmark id in the document."""
    max_id = 0
    for elem in doc.element.body.iter():
        for attr in (qn('w:id'),):
            val = elem.get(attr)
            if val is not None:
                try:
                    max_id = max(max_id, int(val))
                except ValueError:
                    pass
    return max_id

def add_table_captions(doc):
    """
    Walk the body, find each w:tbl, and insert a Caption paragraph after it.
    Skips doc.tables[0] (the empty 1×1 table).
    """
    body = doc.element.body
    bk_id = get_max_bookmark_id(doc) + 1

    table_counter = -1   # -1 so first table (idx 0, empty) gets skipped
    seq_num = 0
    inserted = 0

    children = list(body)
    i = 0
    while i < len(children):
        child = children[i]
        tag = child.tag.split('}')[-1] if '}' in child.tag else child.tag

        if tag == 'tbl':
            table_counter += 1
            if table_counter in TABLE_CAPTIONS:
                seq_num += 1
                caption_text = TABLE_CAPTIONS[table_counter]
                bk_name = f'_Table{seq_num}_Caption'
                cap_p = make_caption_para('Table', seq_num, caption_text, bk_id, bk_name)
                bk_id += 1

                # Insert immediately after the table element
                current_idx = list(body).index(child)
                body.insert(current_idx + 1, cap_p)

                children = list(body)   # refresh after insert
                inserted += 1
                print(f'  [+] Table {seq_num}: "{caption_text}"')

        i += 1

    print(f'[OK] Added {inserted} table captions')
    return seq_num  # total tables captioned

def find_lof_field_end(doc):
    """
    Locate the paragraph(s) that hold the LOF field { TOC \h \z \c "Figure" }.
    Return the last body element that is part of the LOF field block, so we
    can insert the LOT immediately after it.
    """
    body = doc.element.body
    children = list(body)

    W = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'
    lof_end_elem = None
    in_tof_field = False

    for child in children:
        tag = child.tag.split('}')[-1] if '}' in child.tag else child.tag
        if tag != 'p':
            continue

        # Check if this paragraph contains TOC \c "Figure"
        xml_str = etree.tostring(child).decode('utf-8', errors='replace')
        if 'TOC' in xml_str and '"Figure"' in xml_str:
            in_tof_field = True
            lof_end_elem = child

        if in_tof_field:
            lof_end_elem = child
            # Check if field ends in this paragraph
            if 'fldCharType="end"' in xml_str or "fldCharType='end'" in xml_str:
                # This might be the last paragraph of the LOF field block
                # But multi-paragraph LOF: continue until we hit a non-TableofFigures paragraph
                pPr = child.find(f'{{{W}}}pPr')
                if pPr is not None:
                    pStyle = pPr.find(f'{{{W}}}pStyle')
                    style_val = pStyle.get(f'{{{W}}}val', '') if pStyle is not None else ''
                    if style_val not in ('TableofFigures', 'tableoffigures', 'TableOfFigures'):
                        break

    return lof_end_elem

def insert_lot(doc, n_tables):
    """
    Insert a 'List of Tables' heading + TOC \h \z \c "Table" field
    immediately after the List of Figures field block.
    """
    body = doc.element.body
    lof_last = find_lof_field_end(doc)

    if lof_last is None:
        # Fallback: insert before List of Abbreviations heading
        for para in doc.paragraphs:
            if 'List of Abbreviations' in para.text and para.style.name.startswith('Heading'):
                lof_last_idx = list(body).index(para._element) - 1
                lof_last = list(body)[lof_last_idx]
                break

    if lof_last is None:
        print('[WARN] Could not locate insertion point for LOT')
        return

    # Build LOT elements
    lot_heading = make_lot_heading()
    lot_field   = make_toc_field('Table')

    # Find current position and insert in reverse order (so heading comes first)
    children = list(body)
    insert_after_idx = children.index(lof_last)

    # Insert field then heading (reversed because we insert at the same position)
    body.insert(insert_after_idx + 1, lot_field)
    body.insert(insert_after_idx + 1, lot_heading)

    print('[OK] List of Tables heading + TOC field inserted after List of Figures')

# ── Preview helper ────────────────────────────────────────────────────────────

def preview_headings(doc, label='HEADING PREVIEW'):
    """Print a sample of headings with their numPr state."""
    print(f'\n=== {label} ===')
    count = 0
    for para in doc.paragraphs:
        sname = para.style.name
        if sname in ('Heading 3', 'Heading 1', 'Heading 2') and para.text.strip():
            pPr = para._element.find(qn('w:pPr'))
            numPr = pPr.find(qn('w:numPr')) if pPr is not None else None
            if numPr is not None:
                ilvl = numPr.find(qn('w:ilvl'))
                numId = numPr.find(qn('w:numId'))
                iv = ilvl.get(qn('w:val')) if ilvl is not None else '?'
                ni = numId.get(qn('w:val')) if numId is not None else '?'
                print(f'  {sname!r:12s} para numPr(ilvl={iv},numId={ni}) | {para.text.strip()[:60]}')
            else:
                print(f'  {sname!r:12s} para NO-numPr | {para.text.strip()[:60]}')
            count += 1
            if count >= 8:
                break

def preview_style_numpr(doc, label='STYLE numPr CHECK'):
    """Show what numPr the Heading3 style currently carries."""
    print(f'\n=== {label} ===')
    styles_elem = doc.part.styles._element
    for style in styles_elem.findall(qn('w:style')):
        sId = style.get(qn('w:styleId'))
        if sId in ('Heading3', 'Heading2', 'Heading1'):
            pPr = style.find(qn('w:pPr'))
            numPr = pPr.find(qn('w:numPr')) if pPr is not None else None
            if numPr is not None:
                ilvl = numPr.find(qn('w:ilvl'))
                numId = numPr.find(qn('w:numId'))
                iv = ilvl.get(qn('w:val')) if ilvl is not None else '?'
                ni = numId.get(qn('w:val')) if numId is not None else '?'
                print(f'  Style {sId!r}: pPr numPr(ilvl={iv},numId={ni})  <-- DUPLICATE')
            else:
                print(f'  Style {sId!r}: pPr no numPr  [clean]')

# ── Main ──────────────────────────────────────────────────────────────────────

def main():
    shutil.copy2(INPUT, OUTPUT)
    print(f'[OK] Copied original → {OUTPUT}')

    doc = Document(OUTPUT)

    # ── Before preview ────────────────────────────────────────────────────────
    preview_style_numpr(doc, 'BEFORE — Style numPr')
    preview_headings(doc, 'BEFORE — Paragraph numPr (first 8 headings)')

    # ── Fix 1: Remove numPr from Heading3 style ───────────────────────────────
    print('\n--- Applying Fix 1: Remove style-level numPr from Heading3 ---')
    fix_heading3_numpr(doc)

    # ── After preview (Fix 1) ─────────────────────────────────────────────────
    preview_style_numpr(doc, 'AFTER Fix 1 — Style numPr')
    preview_headings(doc, 'AFTER Fix 1 — Paragraph numPr (first 8 headings, UNCHANGED)')

    # ── Fix 2a: Add table captions ────────────────────────────────────────────
    print('\n--- Applying Fix 2a: Add table captions ---')
    n = add_table_captions(doc)

    # ── Fix 2b/c: LOF already correct; insert LOT ────────────────────────────
    print('\n--- Applying Fix 2c: Insert List of Tables after LOF ---')
    insert_lot(doc, n)

    # ── Verify LOF is untouched ───────────────────────────────────────────────
    print('\n=== VERIFICATION: LOF field still present ===')
    for para in doc.paragraphs:
        if para.style.name in ('TableofFigures', 'Table of Figures', 'table of figures'):
            xml_str = etree.tostring(para._element).decode('utf-8', errors='replace')
            if 'TOC' in xml_str and 'Figure' in xml_str:
                print('  [OK] LOF paragraph contains { TOC ... "Figure" } field')
                break

    # ── Verify captions ───────────────────────────────────────────────────────
    print('\n=== VERIFICATION: Caption paragraphs ===')
    for para in doc.paragraphs:
        if para.style.name == 'Caption':
            t = para.text.strip()
            if t:
                print(f'  {t[:80]}')

    # ── Verify LOT ────────────────────────────────────────────────────────────
    print('\n=== VERIFICATION: LOT heading + field ===')
    lot_found = False
    for para in doc.paragraphs:
        if 'List of Tables' in para.text and para.style.name.startswith('Heading'):
            print(f'  Heading: {para.text.strip()!r}')
            lot_found = True
        if lot_found:
            xml_str = etree.tostring(para._element).decode('utf-8', errors='replace')
            if 'TOC' in xml_str and 'Table' in xml_str:
                print('  [OK] LOT paragraph contains { TOC ... "Table" } field')
                break

    # ── Save ──────────────────────────────────────────────────────────────────
    doc.save(OUTPUT)
    print(f'\n[DONE] Saved: {OUTPUT}')
    print(f'       Original unchanged: {INPUT}')

if __name__ == '__main__':
    main()
