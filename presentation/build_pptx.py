# -*- coding: utf-8 -*-
"""
AEGIS FYP deck -> fully editable PowerPoint (python-pptx only).
Every text element is a real PowerPoint text box. Geometry maps the 1920x1080
design 1:1 onto a 13.333in x 7.5in 16:9 slide (EMU/px = 6350). FONT SIZES are
set DIRECTLY in points at a proper projection scale (not px*0.5):
    hero ~80pt · section title 42pt · subhead 27pt · body 21pt (floor 20)
    caption/eyebrow 16pt · footer 14pt
"""
import sys
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

EMU_PX = 6350
def X(px): return Emu(int(round(px * EMU_PX)))

# ── Warm Blossom palette ──────────────────────────────────────────────────
BG      = RGBColor(0xFB,0xF7,0xFF)
INK     = RGBColor(0x2D,0x1A,0x4A)
INK_DIM = RGBColor(0x6E,0x54,0x94)
INK_SOFT= RGBColor(0x9E,0x86,0xC0)
BRAND   = RGBColor(0x7C,0x3A,0xED)
PRIMARY = RGBColor(0xAD,0x90,0xDF)
PRI_DEEP= RGBColor(0x6D,0x44,0xC9)
LINE    = RGBColor(0xEC,0xE3,0xF7)
WHITE   = RGBColor(0xFF,0xFF,0xFF)

DISP = "Sora"        # headings / display (bold via bold flag -> Sora Bold 700)
BODY = "Inter"       # body / labels

EMU_W, EMU_H = 12192000, 6858000

# ── low-level helpers ─────────────────────────────────────────────────────
def set_bg(slide, color):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = color

def _set_alpha(clr_el, alpha):
    clr_el.append(clr_el.makeelement(qn('a:alpha'), {'val': str(int(alpha*100000))}))

def line_alpha(shape, rgb, width_pt, alpha):
    ln = shape.line; ln.color.rgb = rgb; ln.width = Pt(width_pt)
    srgb = ln.color._xFill.find(qn('a:srgbClr'))
    if srgb is not None: _set_alpha(srgb, alpha)

def no_fill(shape): shape.fill.background()
def solid(shape, rgb): shape.fill.solid(); shape.fill.fore_color.rgb = rgb

def fill_alpha(shape, rgb, alpha):
    shape.fill.solid(); shape.fill.fore_color.rgb = rgb
    srgb = shape.fill.fore_color._xFill.find(qn('a:srgbClr'))
    if srgb is not None: _set_alpha(srgb, alpha)

def ring(slide, cx, cy, r, rgb, alpha, w=1.4):
    o = slide.shapes.add_shape(MSO_SHAPE.OVAL, X(cx-r), X(cy-r), X(2*r), X(2*r))
    no_fill(o); o.shadow.inherit = False; line_alpha(o, rgb, w, alpha); return o

def orbit(slide, cx, cy, rx, ry, rot, rgb, alpha, w=1.6):
    o = slide.shapes.add_shape(MSO_SHAPE.OVAL, X(cx-rx), X(cy-ry), X(2*rx), X(2*ry))
    no_fill(o); o.shadow.inherit = False; o.rotation = rot
    line_alpha(o, rgb, w, alpha); return o

def dot(slide, cx, cy, r, rgb, alpha=1.0):
    o = slide.shapes.add_shape(MSO_SHAPE.OVAL, X(cx-r), X(cy-r), X(2*r), X(2*r))
    o.line.fill.background(); o.shadow.inherit = False
    solid(o, rgb) if alpha >= 1 else fill_alpha(o, rgb, alpha); return o

def rect(slide, x, y, w, h, rgb=None, rounded=False):
    shp = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE,
        X(x), X(y), X(w), X(h))
    shp.line.fill.background(); shp.shadow.inherit = False
    solid(shp, rgb) if rgb is not None else no_fill(shp); return shp

def hline(slide, x, y, w, rgb, wt=1.0):
    ln = slide.shapes.add_connector(2, X(x), X(y), X(x+w), X(y))
    ln.line.color.rgb = rgb; ln.line.width = Pt(wt); return ln

def set_tracking(run, pt):
    run._r.get_or_add_rPr().set('spc', str(int(pt*100)))

def outline_text(run, rgb, width_pt, alpha):
    """Hollow/outlined run. a:ln MUST precede the fill in a:rPr."""
    rPr = run._r.get_or_add_rPr()
    for tag in ('a:ln','a:noFill','a:solidFill','a:gradFill','a:blipFill','a:pattFill','a:grpFill'):
        for e in rPr.findall(qn(tag)): rPr.remove(e)
    ln = rPr.makeelement(qn('a:ln'), {'w': str(int(width_pt*12700))})
    sf = ln.makeelement(qn('a:solidFill'), {})
    c = sf.makeelement(qn('a:srgbClr'), {'val': '%02X%02X%02X' % (rgb[0],rgb[1],rgb[2])})
    _set_alpha(c, alpha); sf.append(c); ln.append(sf)
    rPr.insert(0, ln); rPr.insert(1, rPr.makeelement(qn('a:noFill'), {}))

def textbox(slide, x, y, w, h, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(X(x), X(y), X(w), X(h))
    tf = tb.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
    tf.margin_left=0; tf.margin_right=0; tf.margin_top=0; tf.margin_bottom=0
    return tb, tf

def run_in(p, text, font, size_pt, color, bold=False, tracking=None):
    r = p.add_run(); r.text = text; r.font.name = font; r.font.size = Pt(size_pt)
    r.font.bold = bold; r.font.color.rgb = color
    if tracking is not None: set_tracking(r, tracking)
    return r

# ════════════════════════════════════════════════════════════════════════════
def slide_title(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6]); set_bg(s, BG)
    cx, cy = 960, 500
    for r,a in [(250,0.20),(410,0.13),(580,0.08),(770,0.05)]:
        ring(s, cx, cy, r, PRIMARY, a)
    orbit(s, cx, cy, 560, 220, -18, BRAND, 0.12)
    dot(s, 430, 420, 9, BRAND, 0.5); dot(s, 1500, 600, 9, PRIMARY, 0.65)

    chip = rect(s, 700, 56, 520, 96, WHITE, rounded=True)
    chip.line.color.rgb = LINE; chip.line.width = Pt(0.75)
    s.shapes.add_picture("assets/ist-logo.jpeg", X(724), X(74), height=X(60))
    s.shapes.add_picture("assets/aegis-mark.png", X(960-92), X(300), width=X(184))

    tb,tf = textbox(s, 360, 492, 1200, 44, PP_ALIGN.CENTER)
    run_in(tf.paragraphs[0], "FINAL YEAR PROJECT", BODY, 16, BRAND, True, tracking=4)
    tb,tf = textbox(s, 360, 540, 1200, 190, PP_ALIGN.CENTER)
    run_in(tf.paragraphs[0], "AEGIS", DISP, 80, PRI_DEEP, True, tracking=-3)
    tb,tf = textbox(s, 80, 742, 1760, 60, PP_ALIGN.CENTER)
    run_in(tf.paragraphs[0], "Adaptive Edge-Intelligence Guardian and Safety System", DISP, 22, INK, True)
    tb,tf = textbox(s, 310, 802, 1300, 120, PP_ALIGN.CENTER)
    p = tf.paragraphs[0]; p.line_spacing = 1.4
    run_in(p, "On-device stress detection, GPS geofencing and instant SMS alerts.", BODY, 21, INK_DIM)

    hline(s, 120, 926, 1680, LINE)
    def credit(x, w, align, label, lines):
        tb,tf = textbox(s, x, 950, w, 120, align)
        run_in(tf.paragraphs[0], label, BODY, 15, BRAND, True, tracking=1.6)
        for nm, reg in lines:
            pp = tf.add_paragraph(); pp.alignment = align; pp.space_before = Pt(7)
            run_in(pp, nm, DISP, 20, INK, True)
            if reg: run_in(pp, "  "+reg, DISP, 20, INK_SOFT, True)
    credit(120,  600, PP_ALIGN.LEFT,   "PRESENTED BY",
           [("Barira Sarfraz","222201042"),("Nabeeha Zahid","222201009")])
    credit(720,  480, PP_ALIGN.CENTER, "SUPERVISED BY", [("Dr. Altaf Hussain","")])
    credit(1200, 600, PP_ALIGN.RIGHT,  "DEPARTMENT & DATE",
           [("Dept. of Computer Science",""),("KICSIT Campus · June 2026","")])

# ════════════════════════════════════════════════════════════════════════════
FR = [
 ("01","Accounts & Authentication","Firebase Auth login; multiple children."),
 ("02","Live Vitals Dashboard","Real-time HR, SpO2, GSR and temperature."),
 ("03","Location & Geofence Map","Live GPS and safe-zone on OpenStreetMap."),
 ("04","Safe-Zone Configuration","Set the geofence centre and radius."),
 ("05","Alert History","Stress, breach and SpO2 alerts, logged."),
 ("06","Push Notifications","Instant alerts via Firebase Cloud Messaging."),
 ("07","Health Analytics","7-day vitals trends and stress charts."),
 ("08","Calibration Tracking","7-day baseline before alerts activate."),
]

def slide_fr(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6]); set_bg(s, BG)
    cx, cy = 620, 760
    for r,a in [(190,0.24),(340,0.16),(500,0.10),(680,0.06),(870,0.035)]:
        ring(s, cx, cy, r, PRIMARY, a)
    orbit(s, cx, cy, 600, 240, -20, BRAND, 0.10)
    dot(s, 1180, 900, 9, PRIMARY, 0.55); dot(s, 1560, 980, 9, BRAND, 0.40)

    panel = rect(s, 0, 0, 620, 1080)
    try:
        panel.fill.gradient(); panel.fill.gradient_angle = 158
        gs = panel.fill.gradient_stops
        gs[0].color.rgb = RGBColor(0xF1,0xE7,0xFF); gs[0].position = 0.0
        gs[1].color.rgb = RGBColor(0xE9,0xD9,0xFB); gs[1].position = 1.0
    except Exception:
        solid(panel, RGBColor(0xEC,0xDD,0xFB))
    panel.line.fill.background()

    s.shapes.add_picture("assets/aegis-mark.png", X(68), X(78), width=X(96))
    tb,tf = textbox(s, 150, 700, 470, 300)            # ghost numeral
    run_in(tf.paragraphs[0], "08", DISP, 110, BRAND, True)
    outline_text(tf.paragraphs[0].runs[0], (0x7C,0x3A,0xED), 1.2, 0.16)
    tb,tf = textbox(s, 68, 250, 480, 40)
    run_in(tf.paragraphs[0], "REQUIREMENTS", BODY, 16, BRAND, True, tracking=2.5)
    tb,tf = textbox(s, 66, 290, 520, 210)
    p0=tf.paragraphs[0]; p0.line_spacing=1.04
    run_in(p0, "Functional", DISP, 42, INK, True, tracking=-1.5)
    p1=tf.add_paragraph(); p1.line_spacing=1.04
    run_in(p1, "Requirements", DISP, 42, PRI_DEEP, True, tracking=-1.5)
    tb,tf = textbox(s, 68, 482, 470, 170)
    p=tf.paragraphs[0]; p.line_spacing=1.5
    run_in(p, "What the AEGIS parent app delivers across the system.", BODY, 21, INK_DIM)
    rect(s, 68, 956, 200, 52, WHITE, rounded=True)
    tb,tf = textbox(s, 68, 968, 200, 32, PP_ALIGN.CENTER)
    run_in(tf.paragraphs[0], "SECTION 08", BODY, 14, BRAND, True, tracking=1.6)

    col_x = [708, 1314]; col_w = 546
    row_y = [120, 352, 584, 816]
    for i,(n,t,d) in enumerate(FR):
        cxn = col_x[i % 2]; ry = row_y[i // 2]
        tbn,tfn = textbox(s, cxn, ry-2, 84, 64)
        rn = run_in(tfn.paragraphs[0], n, DISP, 30, PRIMARY, True)
        outline_text(rn, (0xAD,0x90,0xDF), 1.2, 1.0)
        tbt,tft = textbox(s, cxn+92, ry-6, col_w-92, 200)
        pt0=tft.paragraphs[0]; pt0.line_spacing=1.08
        run_in(pt0, t, DISP, 27, INK, True, tracking=-1)
        pd=tft.add_paragraph(); pd.line_spacing=1.2; pd.space_before=Pt(7)
        run_in(pd, d, BODY, 21, INK_DIM)
    for r in range(3):
        for cxn in col_x:
            hline(s, cxn, row_y[r]+202, col_w, LINE)

    tb,tf = textbox(s, 708, 1016, 900, 34)
    run_in(tf.paragraphs[0], "Adaptive Edge-Intelligence Guardian and Safety System", BODY, 14, INK_SOFT)
    tb,tf = textbox(s, 1700, 1016, 116, 34, PP_ALIGN.RIGHT)
    run_in(tf.paragraphs[0], "08", DISP, 15, PRI_DEEP, True)

# ════════════════════════════════════════════════════════════════════════════
def main():
    prs = Presentation()
    prs.slide_width = Emu(EMU_W); prs.slide_height = Emu(EMU_H)
    slide_title(prs); slide_fr(prs)
    out = sys.argv[1] if len(sys.argv) > 1 else "aegis_proof.pptx"
    prs.save(out)
    print("saved", out, "- slides:", len(prs.slides._sldIdLst))

if __name__ == "__main__":
    main()
