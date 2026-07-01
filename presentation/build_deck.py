# -*- coding: utf-8 -*-
"""
AEGIS FYP — full editable deck (python-pptx only, no Office automation).
Restructured Descriptive Use Cases (3 stacked Actor/System tables each) and
Detailed Test Cases (labeled cards with PASS/PENDING badges). Content sourced
from the AEGIS thesis docx + project brief in E:\\FypApp. Page numbers are
assigned by slide order automatically.
"""
import sys
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

EMU_PX=6350
def X(px): return Emu(int(round(px*EMU_PX)))

BG=RGBColor(0xFB,0xF7,0xFF); INK=RGBColor(0x2D,0x1A,0x4A); INK_DIM=RGBColor(0x6E,0x54,0x94)
INK_SOFT=RGBColor(0x9E,0x86,0xC0); BRAND=RGBColor(0x7C,0x3A,0xED); PRIMARY=RGBColor(0xAD,0x90,0xDF)
PRI_DEEP=RGBColor(0x6D,0x44,0xC9); PRI_SOFT=RGBColor(0xED,0xE0,0xFF); MIDPAS=RGBColor(0xDD,0xC9,0xF3)
LINE=RGBColor(0xEC,0xE3,0xF7); WHITE=RGBColor(0xFF,0xFF,0xFF); SURF2=RGBColor(0xFA,0xF6,0xFF)
PASS_BG=RGBColor(0xE3,0xF6,0xEC); PASS_FG=RGBColor(0x1E,0x7A,0x4B)
PEND_BG=RGBColor(0xFC,0xEF,0xD6); PEND_FG=RGBColor(0x9A,0x6A,0x12)
FAIL_BG=RGBColor(0xFC,0xE4,0xE8); FAIL_FG=RGBColor(0xC4,0x2A,0x45)
DISP="Sora"; BODY="Inter"
FULLNAME="Adaptive Edge-Intelligence Guardian and Safety System"
EMU_W,EMU_H=12192000,6858000
HERO=68; TITLE=36; SUB=23; BODY_S=19; LEAD=20; CAP=14; EY=14; FOOT=12.5; NUM=26; GHOST=92; BIG=46

def set_bg(s,c): s.background.fill.solid(); s.background.fill.fore_color.rgb=c
def _alpha(el,a): el.append(el.makeelement(qn('a:alpha'),{'val':str(int(a*100000))}))
def _line_a(sh,rgb,w,a):
    ln=sh.line; ln.color.rgb=rgb; ln.width=Pt(w); e=ln.color._xFill.find(qn('a:srgbClr'))
    if e is not None: _alpha(e,a)
def nofill(sh): sh.fill.background()
def solid(sh,rgb): sh.fill.solid(); sh.fill.fore_color.rgb=rgb
def fill_a(sh,rgb,a):
    sh.fill.solid(); sh.fill.fore_color.rgb=rgb; e=sh.fill.fore_color._xFill.find(qn('a:srgbClr'))
    if e is not None: _alpha(e,a)
def ring(s,cx,cy,r,rgb,a,w=1.4):
    o=s.shapes.add_shape(MSO_SHAPE.OVAL,X(cx-r),X(cy-r),X(2*r),X(2*r)); nofill(o); o.shadow.inherit=False; _line_a(o,rgb,w,a)
def orbit(s,cx,cy,rx,ry,rot,rgb,a,w=1.6):
    o=s.shapes.add_shape(MSO_SHAPE.OVAL,X(cx-rx),X(cy-ry),X(2*rx),X(2*ry)); nofill(o); o.shadow.inherit=False; o.rotation=rot; _line_a(o,rgb,w,a)
def dot(s,cx,cy,r,rgb,a=1.0):
    o=s.shapes.add_shape(MSO_SHAPE.OVAL,X(cx-r),X(cy-r),X(2*r),X(2*r)); o.line.fill.background(); o.shadow.inherit=False
    solid(o,rgb) if a>=1 else fill_a(o,rgb,a)
def rect(s,x,y,w,h,rgb=None,rounded=False,line_rgb=None,line_w=0.75):
    sh=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE,X(x),X(y),X(w),X(h)); sh.shadow.inherit=False
    solid(sh,rgb) if rgb is not None else nofill(sh)
    if line_rgb is not None: sh.line.color.rgb=line_rgb; sh.line.width=Pt(line_w)
    else: sh.line.fill.background()
    return sh
def set_dash(sh,val='dash'):
    ln=sh.line._get_or_add_ln(); ln.append(ln.makeelement(qn('a:prstDash'),{'val':val}))
def hline(s,x,y,w,rgb=LINE,wt=1.0):
    ln=s.shapes.add_connector(2,X(x),X(y),X(x+w),X(y)); ln.line.color.rgb=rgb; ln.line.width=Pt(wt); return ln
def track(r,pt): r._r.get_or_add_rPr().set('spc',str(int(pt*100)))
def outline(run,rgb,w,a):
    rPr=run._r.get_or_add_rPr()
    for t in ('a:ln','a:noFill','a:solidFill','a:gradFill','a:blipFill','a:pattFill','a:grpFill'):
        for e in rPr.findall(qn(t)): rPr.remove(e)
    ln=rPr.makeelement(qn('a:ln'),{'w':str(int(w*12700))}); sf=ln.makeelement(qn('a:solidFill'),{})
    c=sf.makeelement(qn('a:srgbClr'),{'val':'%02X%02X%02X'%(rgb[0],rgb[1],rgb[2])}); _alpha(c,a); sf.append(c); ln.append(sf)
    rPr.insert(0,ln); rPr.insert(1,rPr.makeelement(qn('a:noFill'),{}))
def tbox(s,x,y,w,h,align=PP_ALIGN.LEFT,anchor=MSO_ANCHOR.TOP):
    tb=s.shapes.add_textbox(X(x),X(y),X(w),X(h)); tf=tb.text_frame; tf.word_wrap=True; tf.vertical_anchor=anchor
    tf.margin_left=0; tf.margin_right=0; tf.margin_top=0; tf.margin_bottom=0; return tf
def run(p,text,font,sz,color,bold=False,tr=None):
    r=p.add_run(); r.text=text; r.font.name=font; r.font.size=Pt(sz); r.font.bold=bold; r.font.color.rgb=color
    if tr is not None: track(r,tr)
    return r
def T(s,x,y,w,h,text,font,sz,color,bold=False,align=PP_ALIGN.LEFT,ls=None,tr=None,anchor=MSO_ANCHOR.TOP):
    if bold is not None and not isinstance(bold,bool): align=bold; bold=False
    tf=tbox(s,x,y,w,h,align,anchor); p=tf.paragraphs[0]
    if ls: p.line_spacing=ls
    run(p,text,font,sz,color,bold,tr); return tf

def base(prs): s=prs.slides.add_slide(prs.slide_layouts[6]); set_bg(s,BG); return s
def pp(pg): return "%02d"%pg
def panel_rings(s):
    cx,cy=600,770
    for r,a in [(190,0.22),(340,0.15),(500,0.09),(680,0.055),(870,0.032)]: ring(s,cx,cy,r,PRIMARY,a)
    orbit(s,cx,cy,600,240,-20,BRAND,0.09); dot(s,1180,910,8,PRIMARY,0.5); dot(s,1560,980,8,BRAND,0.38)
def center_rings(s):
    cx,cy=960,500
    for r,a in [(250,0.18),(410,0.12),(580,0.075),(770,0.045)]: ring(s,cx,cy,r,PRIMARY,a)
    orbit(s,cx,cy,560,220,-18,BRAND,0.11); dot(s,430,420,8,BRAND,0.45); dot(s,1500,600,8,PRIMARY,0.6)
def light_rings(s):
    for r,a in [(360,0.08),(560,0.05),(820,0.03)]: ring(s,1620,250,r,PRIMARY,a)
def mark(s,x,y,w): s.shapes.add_picture("assets/aegis-mark.png",X(x),X(y),width=X(w))

def panel(s,pg,eyebrow,title_lines,lead=None):
    p=rect(s,0,0,600,1080)
    try:
        p.fill.gradient(); p.fill.gradient_angle=158
        gs=p.fill.gradient_stops; gs[0].color.rgb=RGBColor(0xF1,0xE7,0xFF); gs[0].position=0.0
        gs[1].color.rgb=RGBColor(0xE9,0xD9,0xFB); gs[1].position=1.0
    except Exception: solid(p,RGBColor(0xEC,0xDD,0xFB))
    p.line.fill.background(); mark(s,60,72,84)
    tf=tbox(s,150,724,440,260); run(tf.paragraphs[0],pp(pg),DISP,GHOST,BRAND,True); outline(tf.paragraphs[0].runs[0],(0x7C,0x3A,0xED),1.1,0.15)
    T(s,60,232,470,34,eyebrow.upper(),BODY,EY,BRAND,True,tr=2.4)
    htitle=len(title_lines)*60+16
    tf=tbox(s,58,270,536,htitle)
    for i,(txt,col) in enumerate(title_lines):
        q=tf.paragraphs[0] if i==0 else tf.add_paragraph(); q.line_spacing=1.06
        run(q,txt,DISP,34,col,True,tr=-1)
    if lead:
        ly=270+htitle+18
        T(s,60,ly,468,190,lead,BODY,LEAD,INK_DIM,ls=1.45)
    rect(s,60,956,196,50,WHITE,rounded=True,line_rgb=LINE)
    T(s,60,968,196,30,"SECTION "+pp(pg),BODY,13,BRAND,True,PP_ALIGN.CENTER,tr=1.4)

def footer(s,pg,xl=680):
    T(s,xl,1018,940,30,FULLNAME,BODY,FOOT,INK_SOFT)
    T(s,1740,1016,80,30,pp(pg),DISP,13,PRI_DEEP,True,PP_ALIGN.RIGHT)
def header(s,eyebrow,title,pg,sub=None):
    T(s,120,128,1300,34,eyebrow.upper(),BODY,EY,BRAND,True,tr=2.4)
    T(s,118,166,1560,70,title,DISP,TITLE,INK,True,tr=-1)
    rect(s,122,246,84,5,PRIMARY)
    if sub: T(s,120,262,1500,40,sub,BODY,BODY_S,INK_DIM)
    footer(s,pg,120)
    return 300

# ── content helpers ─────────────────────────────────────────────────────────
def vlist(s,x,y,w,items,row_h):
    for i,(t,d) in enumerate(items):
        ry=y+i*row_h
        tf=tbox(s,x,ry-2,58,50); rr=run(tf.paragraphs[0],"%02d"%(i+1),DISP,NUM,PRIMARY,True); outline(rr,(0xAD,0x90,0xDF),1.1,1.0)
        T(s,x+76,ry-4,w-76,48,t,DISP,SUB,INK,True,tr=-0.5)
        if d: T(s,x+76,ry+SUB*1.7,w-76,row_h-SUB*1.7-8,d,BODY,BODY_S,INK_DIM,ls=1.3)
def card(s,x,y,w,h,title,desc,tag=None):
    rect(s,x,y,w,h,WHITE,rounded=True,line_rgb=LINE)
    T(s,x+28,y+24,w-56,40,title,DISP,SUB,INK,True,tr=-0.5)
    T(s,x+28,y+62,w-56,h-80,desc,BODY,BODY_S,INK_DIM,ls=1.35)
def placeholder(s,x,y,w,h,label,note,tag="DIAGRAM"):
    box=rect(s,x,y,w,h,SURF2,rounded=True,line_rgb=PRIMARY,line_w=2.0); set_dash(box,'dash')
    rect(s,x+w/2-46,y+h/2-120,92,92,PRI_SOFT,rounded=True)
    T(s,x+w/2-46,y+h/2-112,92,72,"+",DISP,52,BRAND,True,PP_ALIGN.CENTER)
    T(s,x+40,y+h/2-8,w-80,52,label,DISP,28,INK,True,PP_ALIGN.CENTER)
    T(s,x+60,y+h/2+44,w-120,80,note,BODY,16,INK_SOFT,PP_ALIGN.CENTER,ls=1.3)
    rect(s,x+w/2-95,y+h-72,190,42,PRI_SOFT,rounded=True)
    T(s,x+w/2-95,y+h-63,190,28,tag+" PLACEHOLDER",BODY,12,BRAND,True,PP_ALIGN.CENTER,tr=1.4)
def phone(s,x,y,w,h,screen,caption):
    rect(s,x,y,w,h,RGBColor(0x2D,0x1A,0x4A),rounded=True)
    rect(s,x+10,y+10,w-20,h-20,SURF2,rounded=True)
    rect(s,x+w/2-32,y+22,64,12,RGBColor(0x2D,0x1A,0x4A),rounded=True)
    T(s,x+16,y+h/2-44,w-32,40,screen,DISP,18,INK,True,PP_ALIGN.CENTER)
    T(s,x+16,y+h/2+2,w-32,40,"App screenshot",BODY,13,INK_SOFT,PP_ALIGN.CENTER)
    T(s,x,y+h+16,w,34,caption,BODY,14,INK_DIM,PP_ALIGN.CENTER)
def cellfmt(cell,text,font,size,color,fill,bold=False,align=PP_ALIGN.LEFT,anchor=MSO_ANCHOR.MIDDLE):
    cell.fill.solid(); cell.fill.fore_color.rgb=fill
    cell.margin_left=X(20); cell.margin_right=X(16); cell.margin_top=X(10); cell.margin_bottom=X(10)
    cell.vertical_anchor=anchor; tf=cell.text_frame; tf.word_wrap=True
    p=tf.paragraphs[0]; p.alignment=align; p.line_spacing=1.15
    run(p,text,font,size,color,bold)
def table_generic(s,x,y,w,col_w,rows,rh=74,hsz=15,bsz=15):
    nr=len(rows); nc=len(rows[0])
    g=s.shapes.add_table(nr,nc,X(x),X(y),X(w),X(rh*nr)); t=g.table; t.first_row=False; t.horz_banding=False
    for ci,cw in enumerate(col_w): t.columns[ci].width=X(cw)
    for ri in range(nr):
        t.rows[ri].height=X(rh)
        for ci in range(nc):
            v=rows[ri][ci]; ish=ri==0
            fill=PRI_SOFT if ish else (WHITE if ri%2 else SURF2)
            col=BRAND if ish else INK
            al=PP_ALIGN.LEFT if ci==0 else (PP_ALIGN.CENTER if nc>2 else PP_ALIGN.LEFT)
            cellfmt(t.cell(ri,ci),v,DISP if ish else BODY,hsz if ish else bsz,col,fill,ish,al)
            if not ish and ci>0:
                rr=t.cell(ri,ci).text_frame.paragraphs[0].runs[0]
                if v in ("Yes","On-device","Free","Offline"): rr.font.color.rgb=PASS_FG; rr.font.bold=True
                elif v in ("No","Cloud","Paid"): rr.font.color.rgb=FAIL_FG; rr.font.bold=True
                elif v=="Limited": rr.font.color.rgb=PEND_FG; rr.font.bold=True
    return t

# ── use-case table (Actor / System), title band + col headers + steps ───────
def uc_table(s,x,y,w,title,colheads,steps,rh_step=78):
    nr=2+len(steps); cw=w//2
    g=s.shapes.add_table(nr,2,X(x),X(y),X(w),X(40+40+rh_step*len(steps))); t=g.table
    t.first_row=False; t.horz_banding=False; t.columns[0].width=X(cw); t.columns[1].width=X(w-cw)
    t.rows[0].height=X(40); a=t.cell(0,0); b=t.cell(0,1); a.merge(b)
    cellfmt(a,title,DISP,17,BRAND,MIDPAS,True,PP_ALIGN.LEFT)
    t.rows[1].height=X(40)
    cellfmt(t.cell(1,0),colheads[0],DISP,15,PRI_DEEP,PRI_SOFT,True)
    cellfmt(t.cell(1,1),colheads[1],DISP,15,PRI_DEEP,PRI_SOFT,True)
    for i,(L,R) in enumerate(steps):
        t.rows[2+i].height=X(rh_step); fill=WHITE if i%2==0 else SURF2
        cellfmt(t.cell(2+i,0),L,BODY,16,INK,fill,False,PP_ALIGN.LEFT,MSO_ANCHOR.TOP)
        cellfmt(t.cell(2+i,1),R,BODY,16,INK_DIM,fill,False,PP_ALIGN.LEFT,MSO_ANCHOR.TOP)
    return y+40+40+rh_step*len(steps)

def uc_header(s,pg,no,title,cont=False):
    eb="Use Case %02d"%no + (" (continued)" if cont else "")
    header(s,eb,title,pg)

def badge(s,x,y,status):
    bg,fg={'PASS':(PASS_BG,PASS_FG),'PENDING':(PEND_BG,PEND_FG),'FAIL':(FAIL_BG,FAIL_FG)}[status]
    rect(s,x,y,150,42,bg,rounded=True); T(s,x,y+9,150,26,status,DISP,14,fg,True,PP_ALIGN.CENTER,tr=1.0)

def test_card(s,x,y,w,h,tcid,title,obj,inp,exp,act,crit,status):
    rect(s,x,y,w,h,WHITE,rounded=True,line_rgb=LINE)
    rect(s,x,y+12,8,h-24,PRIMARY,rounded=True)            # pastel left border
    ix=x+40; iw=w-80
    T(s,ix,y+26,iw-170,26,tcid,BODY,15,BRAND,True,tr=1.4)
    T(s,ix,y+52,iw-170,66,title,DISP,22,INK,True,ls=1.05,tr=-0.5)
    badge(s,x+w-186,y+30,status)
    hline(s,ix,y+130,iw)
    rows=[("OBJECTIVE",obj),("INPUT",inp),("EXPECTED RESULT",exp),("ACTUAL RESULT",act),("PASS CRITERIA",crit)]
    ry=y+152
    for lab,val in rows:
        T(s,ix,ry,iw,22,lab,BODY,13,PRI_DEEP,True,tr=1.6)
        T(s,ix,ry+24,iw,44,val,BODY,18,INK_DIM,ls=1.2)
        ry+=86

# ════════════════════════════════════════════════════════════════════════════
def s_title(prs,pg):
    s=base(prs); center_rings(s)
    rect(s,700,56,520,96,WHITE,rounded=True,line_rgb=LINE); s.shapes.add_picture("assets/ist-logo.jpeg",X(724),X(74),height=X(60))
    mark(s,960-86,300,172)
    T(s,360,498,1200,40,"FINAL YEAR PROJECT",BODY,15,BRAND,True,PP_ALIGN.CENTER,tr=4)
    T(s,360,540,1200,148,"AEGIS",DISP,HERO,PRI_DEEP,True,PP_ALIGN.CENTER,tr=-2)
    T(s,80,702,1760,56,FULLNAME,DISP,22,INK,True,PP_ALIGN.CENTER)
    T(s,360,762,1200,90,"A smart child-safety wearable with on-device stress detection, GPS geofencing and instant SMS alerts.",BODY,18,INK_DIM,PP_ALIGN.CENTER,ls=1.4)
    hline(s,120,890,1680)
    def credit(x,w,al,label,lines):
        T(s,x,912,w,30,label,BODY,14,BRAND,True,al,tr=1.6)
        tf=tbox(s,x,946,w,120,al)
        for j,(nm,reg) in enumerate(lines):
            q=tf.paragraphs[0] if j==0 else tf.add_paragraph(); q.alignment=al; q.space_before=Pt(6)
            run(q,nm,DISP,19,INK,True)
            if reg: run(q,"  "+reg,DISP,19,INK_SOFT,True)
    credit(120,620,PP_ALIGN.LEFT,"PRESENTED BY",[("Barira Sarfraz","222201042"),("Nabeeha Zahid","222201009")])
    credit(740,440,PP_ALIGN.CENTER,"SUPERVISED BY",[("Dr. Altaf Hussain","")])
    credit(1200,600,PP_ALIGN.RIGHT,"DEPARTMENT & DATE",[("Dept. of Computer Science",""),("KICSIT Campus · June 2026","")])

def s_agenda(prs,pg):
    s=base(prs); panel_rings(s); top=header(s,"Outline","Agenda",pg)
    items=[("Introduction & Background","The problem and why AEGIS exists"),
           ("Objectives, Scope & Methodology","Goals and how we built it"),
           ("Literature & Comparison","Prior work and where AEGIS stands"),
           ("Solution & Requirements","Our system, functional & non-functional"),
           ("System Design & Diagrams","Architecture, hardware, UML"),
           ("Dataset, ML Model & Results","WESAD, Random Forest, accuracy"),
           ("Testing & Mobile App","Use cases, test cases and the app"),
           ("Contribution & Conclusion","Team, limitations, future work")]
    for c in range(2):
        for r in range(4):
            i=c*4+r; x=120+c*860; y=top+18+r*168
            tf=tbox(s,x,y-2,58,50); rr=run(tf.paragraphs[0],"%02d"%(i+1),DISP,NUM,PRIMARY,True); outline(rr,(0xAD,0x90,0xDF),1.1,1.0)
            T(s,x+74,y-4,746,44,items[i][0],DISP,SUB,INK,True,tr=-0.5)
            T(s,x+74,y+38,746,60,items[i][1],BODY,BODY_S,INK_DIM,ls=1.3)

def s_intro(prs,pg):
    s=base(prs); panel_rings(s); panel(s,pg,"Introduction",[("Introduction",INK),("& Background",PRI_DEEP)]); footer(s,pg)
    T(s,680,150,1140,130,"Every year, parents of young children face the same fear: a child who wanders off, panics in a crowd, or faces a medical scare cannot always say so in time.",BODY,LEAD,INK_DIM,ls=1.5)
    rect(s,680,312,1140,2,LINE)
    T(s,680,340,1140,40,"The problem",DISP,SUB,PRI_DEEP,True,tr=-0.5)
    T(s,680,388,1140,150,"Existing child trackers only report GPS location and depend on constant internet. They cannot sense distress, they go silent where there is no signal, and they say nothing about how a child actually feels.",BODY,BODY_S,INK_DIM,ls=1.45)
    T(s,680,556,1140,40,"Our answer",DISP,SUB,PRI_DEEP,True,tr=-0.5)
    T(s,680,604,1140,170,"AEGIS is a wearable band for children aged 4-12 that reads physiological signals in real time, detects stress or danger on the device itself, tracks location, and alerts parents over SMS - working even with no internet.",BODY,BODY_S,INK_DIM,ls=1.45)

def s_objectives(prs,pg):
    s=base(prs); panel_rings(s); panel(s,pg,"Goals",[("Objectives",INK),("& Scope",PRI_DEEP)],
        "What AEGIS sets out to achieve, and the boundaries we set."); footer(s,pg)
    vlist(s,680,150,1150,[
        ("Sense distress, not just location","Read heart rate, SpO2, skin response and temperature on the wrist."),
        ("Decide on the edge","Run the stress model on the ESP32 so alerts work without the cloud."),
        ("Alert reliably anywhere","Send SMS over GSM so parents are reached even with no internet."),
        ("Keep every child personal","Calibrate a 7-day baseline so detection fits the individual child."),
        ("Stay private and free","Restrict data per parent and build only on free-tier services.")],164)

def s_methodology(prs,pg):
    s=base(prs); panel_rings(s); panel(s,pg,"Approach",[("Methodology",INK),("& Development",PRI_DEEP)],
        "An incremental, hardware-independent build so progress never stalls."); footer(s,pg)
    vlist(s,680,150,1150,[
        ("Design first","Architecture, data model and UML defined before code."),
        ("App & backend in parallel","Flutter screens built with dummy data, then wired to Firebase."),
        ("Model offline","Random Forest trained and validated on the WESAD dataset."),
        ("Firmware as pure logic","Z-score, sliding-window and geofence written and unit-tested off-device."),
        ("Integrate & test","Hardware, model, app and alerts brought together for end-to-end testing.")],158)

def s_literature(prs,pg):
    s=base(prs); panel_rings(s); top=header(s,"Background","Literature Review",pg)
    cards=[("Wearable stress detection","The WESAD dataset (Schmidt et al., 2018) established wrist signals - HR, EDA, temperature, motion - as reliable cues for stress versus calm states."),
           ("On-device / edge ML","Research on TinyML shows compact models such as Random Forests can run on microcontrollers, enabling decisions without a network round-trip."),
           ("Child-tracking devices","Commercial trackers focus on GPS and calling; studies note their dependence on connectivity and the absence of physiological awareness."),
           ("Physiological signals","Galvanic skin response and heart-rate variability are documented markers of acute stress, motivating multi-sensor fusion over location alone.")]
    for i,(t,d) in enumerate(cards):
        x=120+(i%2)*880; y=top+12+(i//2)*232; card(s,x,y,830,208,t,d)

def s_comparison(prs,pg):
    s=base(prs); panel_rings(s); top=header(s,"Where We Stand","AEGIS vs Existing Wearables",pg)
    rows=[["Capability","AEGIS","GPS Tracker","Kids' Smartwatch"],
          ["Live GPS location","Yes","Yes","Yes"],
          ["Physiological / stress sensing","Yes","No","Limited"],
          ["On-device (edge) decisions","On-device","Cloud","Cloud"],
          ["Works with no internet","Offline","No","No"],
          ["SMS alert fallback","Yes","No","Limited"],
          ["Geofencing","Yes","Limited","Yes"],
          ["Backend cost","Free","Paid","Paid"]]
    table_generic(s,120,top+6,1680,[560,360,380,380],rows,rh=84,bsz=17)

def s_solution(prs,pg):
    s=base(prs); panel_rings(s); top=header(s,"The System","Our Solution",pg)
    comp=[("01  Wearable Band","ESP32 reads the sensors, runs the model, calibrates and decides - the brain of the system."),
          ("02  Edge ML Model","A Random Forest trained on WESAD, flashed to the band, classifying stress entirely on-device."),
          ("03  Parent App","A Flutter app that visualises vitals, location, alerts and analytics - it never makes decisions."),
          ("04  Firebase Backend","Free-tier Firestore, Auth and Cloud Messaging for storage, sync and push.")]
    for i,(t,d) in enumerate(comp):
        x=120+(i%2)*880; y=top+12+(i//2)*228; card(s,x,y,830,206,t,d)

def s_fr(prs,pg):
    s=base(prs); panel_rings(s); panel(s,pg,"Requirements",[("Functional",INK),("Requirements",PRI_DEEP)]); footer(s,pg)
    FR=[("Accounts & Authentication","Firebase Auth login; multiple children."),
        ("Live Vitals Dashboard","Real-time HR, SpO2, GSR and temperature."),
        ("Location & Geofence Map","Live GPS and safe-zone on OpenStreetMap."),
        ("Safe-Zone Configuration","Set the geofence centre and radius."),
        ("Alert History","Stress, breach and SpO2 alerts, logged."),
        ("Push Notifications","Instant alerts via Firebase Cloud Messaging."),
        ("Health Analytics","7-day vitals trends and stress charts."),
        ("Calibration Tracking","7-day baseline before alerts activate.")]
    for i,(t,d) in enumerate(FR):
        x=680+(i%2)*585; y=150+(i//2)*215
        tf=tbox(s,x,y-2,58,50); rr=run(tf.paragraphs[0],"%02d"%(i+1),DISP,24,PRIMARY,True); outline(rr,(0xAD,0x90,0xDF),1.1,1.0)
        T(s,x+70,y-4,505,46,t,DISP,21,INK,True,tr=-0.5)
        T(s,x+70,y+38,505,120,d,BODY,BODY_S,INK_DIM,ls=1.3)

def s_nfr(prs,pg):
    s=base(prs); panel_rings(s); panel(s,pg,"Quality",[("Non-Functional",INK),("Requirements",PRI_DEEP)]); footer(s,pg)
    NF=[("Reliability","Alerts fire even offline - SMS over GSM, cached then synced."),
        ("Security & Privacy","Firestore rules restrict each parent to their own children's data."),
        ("Performance","A 10-second monitoring cycle with real-time app updates."),
        ("Power Efficiency","Duty-cycled sensing to extend battery life on a small LiPo cell."),
        ("Usability","The 7-day calibration is communicated clearly to parents."),
        ("Affordability","Built entirely on free-tier services - zero paid APIs.")]
    for i,(t,d) in enumerate(NF):
        x=680+(i%2)*585; y=150+(i//2)*278
        dot(s,x+8,y+14,8,PRIMARY)
        T(s,x+34,y-4,540,44,t,DISP,SUB,INK,True,tr=-0.5)
        T(s,x+34,y+42,540,180,d,BODY,BODY_S,INK_DIM,ls=1.35)

def s_tools(prs,pg):
    s=base(prs); panel_rings(s); top=header(s,"Stack","Tools & Technologies",pg)
    groups=[("Wearable","ESP32 · MAX30102 · MPU6050 · GSR · MLX90614 · NEO-6M GPS · SIM800L GSM"),
            ("Machine Learning","Python · scikit-learn · Random Forest · WESAD dataset"),
            ("Mobile App","Flutter · Dart · Provider · fl_chart · flutter_map"),
            ("Backend (free tier)","Firebase Auth · Cloud Firestore · Cloud Messaging"),
            ("Maps & Comms","OpenStreetMap · WiFi · BLE · GSM / SMS"),
            ("Tooling","GitHub · GitHub Actions CI · PlatformIO")]
    for i,(t,d) in enumerate(groups):
        x=120+(i%3)*565; y=top+12+(i//3)*250; card(s,x,y,530,224,t,d)

def diagram_slide(prs,pg,eyebrow,title_lines,label,note,tag="DIAGRAM"):
    s=base(prs); panel_rings(s); panel(s,pg,eyebrow,title_lines); footer(s,pg)
    placeholder(s,680,150,1140,820,label,note,tag)

# Use case data: (no,title, typical, alt, exception)
UCS=[
 (1,"View Real-Time Vitals",
  [("1. Parent logs into the app.","1. System authenticates the credentials and grants access."),
   ("2. Parent selects the child profile.","1. System retrieves the data for the selected child's band."),
   ("3. Parent opens the vitals dashboard.","1. System fetches the latest HR, SpO2, temperature and stress. 2. Displays them on the dashboard in real time.")],
  [("1a. Parent opens the app while offline.","1a. System displays the last cached reading from the band."),
   ("1b. Connectivity is restored.","1b. System resumes the typical flow with live data.")],
  [("1c. Parent enters incorrect login credentials.","1c. System prompts re-entry of credentials."),
   ("1d. Login fails three times.","1d. System locks the account for five minutes.")]),
 (2,"Monitor Heart Rate & SpO2",
  [("1. MAX30102 reads pulse and oxygen data.","1. ESP32 computes heart rate and SpO2 values."),
   ("2. ESP32 compares readings to baseline thresholds.","1. System uploads the validated readings to the cloud."),
   ("3. Parent opens the app to check trends.","1. App fetches and displays HR and SpO2 history.")],
  [("1a. Sensor noise is detected during reading.","1a. Band filters the signal with the onboard ML model, then resumes once clean.")],
  [("1c. MAX30102 fails to respond.","1c. Band logs the sensor error and sends a fallback alert via GSM.")]),
 (3,"Monitor Body Temperature",
  [("1. MLX90614 reads the child's body temperature.","1. ESP32 compares the value to the configured threshold."),
   ("2. Reading is within the normal range.","1. System stores the reading in the cloud."),
   ("3. Parent views the temperature trend.","1. App displays the latest and historical temperature data.")],
  [("1a. Ambient interference affects the reading.","1a. Band recalibrates the sensor and resumes with a corrected reading.")],
  [("1c. Reading is out of range (>50C or <20C).","1c. System activates anomaly detection and triggers an alert.")]),
 (4,"Monitor Stress Level (GSR)",
  [("1. GSR sensor measures skin conductance.","1. ESP32 runs the on-device ML model to infer stress."),
   ("2. Stress level is classified.","1. System streams the classified data to the cloud."),
   ("3. Parent opens the app.","1. App displays the current and historical stress trend.")],
  [("1a. GSR signal is weak or unstable.","1a. System applies the child's baseline, then resumes.")],
  [("1c. Stress remains persistently high.","1c. System triggers an alert without waiting for cloud confirmation.")]),
 (5,"View Live Location & History",
  [("1. Parent opens the location tab.","1. App requests the latest location from the cloud."),
   ("2. NEO-6M GPS provides coordinates.","1. ESP32 checks the coordinates against the geofence."),
   ("3. Band reports the verified position.","1. System uploads coordinates and updates the map view.")],
  [("1a. GPS signal is weak or unavailable.","1a. Band falls back to cell-tower location and resumes.")],
  [("1c. Location upload fails.","1c. System stores data locally and retries once reconnected.")]),
]
def uc_one(prs,pg,uc):                 # all three tables on one slide
    no,title,typ,alt,exc=uc
    s=base(prs); light_rings(s); uc_header(s,pg,no,title)
    yb=uc_table(s,120,278,1680,"Typical Course of Actions",("Actor's Action","System Response"),typ)
    yb=uc_table(s,120,yb+22,1680,"Alternative Course of Actions",("Actor's Action","Alternative System Response"),alt)
    uc_table(s,120,yb+22,1680,"Exception Flow",("Actor's Action","Exception System Response"),exc)
def uc_typ(prs,pg,uc):                 # split A: typical only
    no,title,typ,alt,exc=uc
    s=base(prs); light_rings(s); uc_header(s,pg,no,title)
    uc_table(s,120,300,1680,"Typical Course of Actions",("Actor's Action","System Response"),typ)
def uc_altexc(prs,pg,uc):              # split B: alternative + exception
    no,title,typ,alt,exc=uc
    s=base(prs); light_rings(s); uc_header(s,pg,no,title,cont=True)
    yb=uc_table(s,120,300,1680,"Alternative Course of Actions",("Actor's Action","Alternative System Response"),alt)
    uc_table(s,120,yb+28,1680,"Exception Flow",("Actor's Action","Exception System Response"),exc)

def s_dataset(prs,pg):
    s=base(prs); panel_rings(s); panel(s,pg,"Learning",[("Dataset &",INK),("ML Model",PRI_DEEP)],
        "Trained offline, then flashed to the band to run on-device."); footer(s,pg)
    rows=[("Dataset","WESAD - 15 subjects, wrist physiological signals"),
          ("Algorithm","Random Forest · 15 trees · max depth 8"),
          ("Output","Binary: NORMAL or STRESS"),
          ("Features","14 inputs - HR, HRV, RMSSD, EDA, temperature, motion"),
          ("On-device","Exported to a C header and run on the ESP32")]
    for i,(k,v) in enumerate(rows):
        y=150+i*160
        T(s,680,y,360,40,k,DISP,SUB,PRI_DEEP,True,tr=-0.5)
        T(s,680,y+46,1140,90,v,BODY,BODY_S,INK_DIM,ls=1.35)
        if i<len(rows)-1: hline(s,680,y+138,1140)

def s_mlresults(prs,pg):
    s=base(prs); panel_rings(s); top=header(s,"Performance","ML Model Results",pg)
    for i,(big,lab) in enumerate([("97.96%","Training accuracy"),("96.03%","Cross-validation"),("± 1.24%","CV std. deviation")]):
        x=120+i*410; y=top+12; rect(s,x,y,380,196,WHITE,rounded=True,line_rgb=LINE)
        T(s,x+26,y+34,332,70,big,DISP,BIG,BRAND,True); T(s,x+26,y+116,332,50,lab,BODY,BODY_S,INK_DIM)
    placeholder(s,120,top+236,750,468,"Confusion Matrix","NORMAL vs STRESS - export from ml/test_aegis_model.py","CHART")
    rect(s,900,top+236,900,468,SURF2,rounded=True,line_rgb=LINE)
    T(s,936,top+274,820,40,"Per-class metrics",DISP,SUB,INK,True,tr=-0.5)
    T(s,936,top+324,820,110,"Precision, recall and F1 (per NORMAL / STRESS class) come from the model test script's classification report and will be filled in here.",BODY,BODY_S,INK_DIM,ls=1.4)
    for j,m in enumerate(["Precision  —","Recall  —","F1-score  —"]):
        T(s,936,top+452+j*66,820,40,m,DISP,20,PRI_DEEP,True)

# Test cards: (id,title,objective,input,expected,actual,criteria,status)
TCARDS=[
 ("ML-EVAL-01","Model Classification Accuracy","Evaluate the RF classifier on WESAD.","WESAD: 14 features, 15 subjects.","Accuracy at or above 90%.","97.96% train; 96.03% CV (+/-1.24%).","Accuracy >= 90%, stable across folds.","PASS"),
 ("TC-FW-001","Baseline Z-Score Normalization","Standardize a reading vs the baseline.","reading=80, mean=70, std=5.","Returns 2.0.","Pending execution (MVP).","Z-score equals the computed value.","PENDING"),
 ("TC-FW-002","Stress Classification","Confirm the RF returns the expected label.","A STRESS vector, then a NORMAL vector.","Returns STRESS, then NORMAL.","Pending execution (MVP).","Labels match training annotations.","PENDING"),
 ("TC-FW-003","7-Day Calibration Window","Finish calibration only after 7 days.","7 days of mock HR/EDA/temp readings.","Classification enabled from day 8.","Pending execution (MVP).","Baseline finalized; gated to day 8.","PENDING"),
 ("TC-HW-001","Heart Rate Accuracy vs Oximeter","MAX30102 HR vs a clinical oximeter.","10 paired samples over 5 minutes.","Within +/- 5 bpm of reference.","Pending hardware session.","Mean abs. difference <= 5 bpm.","PENDING"),
 ("TC-HW-002","SMS Fallback When WiFi Down","GSM SMS fallback when WiFi fails.","Stress alert with WiFi disabled.","SMS received within fallback window.","Pending hardware session.","Logs 'SMS sent: OK'; SMS received.","PENDING"),
 ("TC-APP-001","Parent Account Registration","Register a parent via Firebase Auth.","Name, unique email, valid password.","Account created; onboarding opens.","Pending execution (MVP).","Account in Firebase Auth; navigates.","PENDING"),
 ("TC-APP-002","Define and Persist a Safe Zone","Draw a safe zone; check it persists.","Zone name, centre, radius.","Zone reappears after restart.","Pending execution (MVP).","Geofence saved in Firestore.","PENDING"),
 ("TC-APP-003","Invalid Login Rejection","Reject wrong credentials clearly.","Valid email, wrong password.","Error shown; no navigation.","Pending execution (MVP).","Stays on login; no token issued.","PENDING"),
 ("TC-INT-001","End-to-End Geofence Breach","Breach: firmware through to app push.","Mock GPS outside the safe zone.","Push arrives within sync interval.","Pending integration session.","Alert in Firestore; push delivered.","PENDING"),
 ("TC-INT-002","Alert Flag Triggers FCM Push","Alert write triggers an FCM push.","Vitals record with alert = true.","Push received; sync unaffected.","Pending integration session.","Delivery log shows dispatch.","PENDING"),
]
def tc_slide(prs,pg,cards):
    s=base(prs); light_rings(s); header(s,"Verification","Detailed Test Cases",pg)
    if len(cards)==2:
        test_card(s,120,290,810,600,*cards[0]); test_card(s,990,290,810,600,*cards[1])
    else:
        test_card(s,555,290,810,600,*cards[0])

def s_appdesign(prs,pg):
    s=base(prs); panel_rings(s); header(s,"Product","Mobile App Design",pg,
        "The parent app's real flow - drop native captures into the frames.")
    screens=[("Splash & Login","Sign in"),("Dashboard","Live vitals"),("Map & Geofence","Location"),("Alerts","History"),("Analytics","Trends")]
    fw,fh=300,560; gap=(1920-240-5*fw)/4
    for i,(scr,cap) in enumerate(screens):
        phone(s,120+i*(fw+gap),360,fw,fh,scr,cap)

def s_team(prs,pg):
    s=base(prs); panel_rings(s); panel(s,pg,"Credits",[("Team",INK),("Contribution",PRI_DEEP)],
        "A shared build - refine the split to match your own division of work."); footer(s,pg)
    people=[("Barira Sarfraz","222201042",["Flutter app screens & UI","Firebase integration","App–backend data flow"]),
            ("Nabeeha Zahid","222201009",["ML model & WESAD pipeline","Firmware logic & geofencing","Documentation & testing"])]
    for i,(nm,reg,tasks) in enumerate(people):
        x=680+i*585; rect(s,x,150,545,720,WHITE,rounded=True,line_rgb=LINE); mark(s,x+34,188,64)
        T(s,x+112,196,400,40,nm,DISP,SUB,INK,True,tr=-0.5); T(s,x+112,242,400,30,reg,BODY,15,INK_SOFT)
        hline(s,x+34,308,477)
        for j,tk in enumerate(tasks):
            yy=348+j*108; dot(s,x+44,yy+12,7,PRIMARY); T(s,x+70,yy,440,90,tk,BODY,BODY_S,INK_DIM,ls=1.3)

def s_challenges(prs,pg):
    s=base(prs); panel_rings(s); top=header(s,"Problem Solving","Challenges & Solutions",pg)
    ch=[("Hardware not yet built","Wrote sensor logic as pure, testable functions and simulated data so app and model could progress."),
        ("No reliable internet","Moved all decisions on-device and added GSM SMS with offline caching that syncs later."),
        ("Stress vs. physical play","Added an accelerometer activity filter so high motion with low skin response is not flagged."),
        ("Every child is different","A 7-day calibration builds a personal baseline; readings become Z-scores against it."),
        ("Sensitive children's data","Firestore security rules scope every record to the owning parent only."),
        ("Zero budget","Chose a fully free-tier stack - Firebase, OpenStreetMap, GSM hardware only.")]
    for i,(t,d) in enumerate(ch):
        x=120+(i%2)*880; y=top+10+(i//2)*230; rect(s,x,y,830,208,WHITE,rounded=True,line_rgb=LINE)
        T(s,x+28,y+22,774,40,t,DISP,21,INK,True,tr=-0.5); T(s,x+28,y+66,774,120,d,BODY,BODY_S,INK_DIM,ls=1.35)

def s_limits(prs,pg):
    s=base(prs); panel_rings(s); panel(s,pg,"Honesty",[("Limitations &",INK),("Future Work",PRI_DEEP)]); footer(s,pg)
    T(s,680,150,540,40,"Where we are today",DISP,SUB,PRI_DEEP,True,tr=-0.5)
    for j,t in enumerate(["Hardware is at dev-kit prototype stage","Some app screens use simulated data","The model classifies stress as binary only","Battery life is not yet measured"]):
        yy=206+j*116; dot(s,687,yy+12,7,PRIMARY); T(s,714,yy,500,100,t,BODY,BODY_S,INK_DIM,ls=1.3)
    T(s,1270,150,540,40,"Where we are headed",DISP,SUB,PRI_DEEP,True,tr=-0.5)
    for j,t in enumerate(["Full sensor and band integration","BLE pairing and power optimization","More affect classes beyond stress","A real-world pilot with families"]):
        yy=206+j*116; dot(s,1277,yy+12,7,PRIMARY); T(s,1304,yy,500,100,t,BODY,BODY_S,INK_DIM,ls=1.3)

def s_poster(prs,pg):
    s=base(prs); panel_rings(s); panel(s,pg,"Showcase",[("Project",INK),("Poster",PRI_DEEP)],
        "The A1 project poster summarising AEGIS for the exhibition."); footer(s,pg)
    placeholder(s,680,150,1140,820,"Project Poster","Drop the final exhibition poster (portrait A1) here","POSTER")

def s_conclusion(prs,pg):
    s=base(prs); panel_rings(s); panel(s,pg,"Closing",[("Conclusion",PRI_DEEP)]); footer(s,pg)
    T(s,680,160,1140,200,"AEGIS reimagines child safety as more than a dot on a map. By sensing distress on the wrist, deciding on the device, and reaching parents even without internet, it turns a tracker into a true guardian.",BODY,LEAD,INK_DIM,ls=1.5)
    rect(s,680,386,1140,2,LINE)
    take=[("Edge-first","Decisions happen on the band, not the cloud."),
          ("Always reachable","SMS over GSM means alerts arrive offline."),
          ("Personal & private","Per-child calibration, per-parent data."),
          ("Built for zero cost","A complete system on free-tier services.")]
    for i,(t,d) in enumerate(take):
        x=680+(i%2)*585; y=424+(i//2)*228
        T(s,x,y,545,40,t,DISP,SUB,INK,True,tr=-0.5); T(s,x,y+46,545,150,d,BODY,BODY_S,INK_DIM,ls=1.35)

def s_references(prs,pg):
    s=base(prs); panel_rings(s); top=header(s,"Sources","References",pg)
    refs=["Schmidt, P. et al. (2018). Introducing WESAD, a Multimodal Dataset for Wearable Stress and Affect Detection. ACM ICMI.",
          "Breiman, L. (2001). Random Forests. Machine Learning, 45(1), 5-32.",
          "Pedregosa, F. et al. (2011). Scikit-learn: Machine Learning in Python. JMLR 12.",
          "Flutter Documentation. Google. flutter.dev",
          "Firebase Documentation (Authentication, Firestore, Cloud Messaging). Google.",
          "OpenStreetMap Contributors. openstreetmap.org",
          "Espressif Systems. ESP32 Technical Reference Manual.",
          "fl_chart & flutter_map - Flutter community packages. pub.dev"]
    for i,r in enumerate(refs):
        x=120+(i%2)*880; y=top+16+(i//2)*168
        tf=tbox(s,x,y,40,40); run(tf.paragraphs[0],str(i+1),DISP,18,PRIMARY,True)
        T(s,x+44,y,786,150,r,BODY,17,INK_DIM,ls=1.35)

def s_thanks(prs,pg):
    s=base(prs); center_rings(s); mark(s,960-80,318,160)
    T(s,360,520,1200,90,"Thank You",DISP,64,PRI_DEEP,True,PP_ALIGN.CENTER,tr=-2)
    T(s,360,636,1200,50,"Questions & Discussion",BODY,22,INK_DIM,PP_ALIGN.CENTER)
    hline(s,660,724,600)
    T(s,360,752,1200,40,"AEGIS · "+FULLNAME,BODY,16,INK_SOFT,PP_ALIGN.CENTER)
    T(s,360,792,1200,40,"Barira Sarfraz · Nabeeha Zahid    |    Supervisor: Dr. Altaf Hussain",BODY,15,INK_SOFT,PP_ALIGN.CENTER)

# ════════════════════════════════════════════════════════════════════════════
def main():
    prs=Presentation(); prs.slide_width=Emu(EMU_W); prs.slide_height=Emu(EMU_H)
    B=[
      s_title, s_agenda, s_intro, s_objectives, s_methodology, s_literature, s_comparison,
      s_solution, s_fr, s_nfr, s_tools,
      lambda p,n: diagram_slide(p,n,"Architecture",[("System",INK),("Architecture",PRI_DEEP)],"System Architecture","Four layers: Wearable band → GSM / WiFi → Firebase → Parent app"),
      lambda p,n: diagram_slide(p,n,"Data Flow",[("System Workflow",INK),("& Data Flow",PRI_DEEP)],"Workflow / Data Flow","Sensor read → Z-score → ML → sliding window → alert → SMS / Firebase"),
      lambda p,n: diagram_slide(p,n,"Hardware",[("Hardware Block",INK),("& Circuit",PRI_DEEP)],"Hardware Block / Circuit","ESP32 with MAX30102, MPU6050, GSR, MLX90614, NEO-6M and SIM800L"),
      lambda p,n: diagram_slide(p,n,"Behaviour",[("Use Case",INK),("Diagram",PRI_DEEP)],"Use Case Diagram","Parent and Wearable actors across monitoring, alerts and configuration"),
      lambda p,n: uc_typ(p,n,UCS[0]),
      lambda p,n: uc_altexc(p,n,UCS[0]),
      lambda p,n: uc_one(p,n,UCS[1]),
      lambda p,n: uc_one(p,n,UCS[2]),
      lambda p,n: uc_one(p,n,UCS[3]),
      lambda p,n: uc_one(p,n,UCS[4]),
      lambda p,n: diagram_slide(p,n,"Data Model",[("Entity-Relationship",INK),("Diagram",PRI_DEEP)],"ERD","users → children → vitals / alerts, and devices"),
      lambda p,n: diagram_slide(p,n,"Behaviour",[("Activity",INK),("Diagram",PRI_DEEP)],"Activity Diagram","Monitoring cycle: calibrate → monitor → detect → alert"),
      lambda p,n: diagram_slide(p,n,"Structure",[("Class",INK),("Diagram",PRI_DEEP)],"Class Diagram","Models (Child, Vital, Alert) and services (Auth, Firestore, Map)"),
      lambda p,n: diagram_slide(p,n,"Interaction",[("Sequence",INK),("Diagram",PRI_DEEP)],"Sequence Diagram","Breach alert across Wearable → GSM → Parent and Wearable → Firebase → App"),
      s_dataset, s_mlresults,
      lambda p,n: tc_slide(p,n,[TCARDS[0],TCARDS[1]]),
      lambda p,n: tc_slide(p,n,[TCARDS[2],TCARDS[3]]),
      lambda p,n: tc_slide(p,n,[TCARDS[4],TCARDS[5]]),
      lambda p,n: tc_slide(p,n,[TCARDS[6],TCARDS[7]]),
      lambda p,n: tc_slide(p,n,[TCARDS[8],TCARDS[9]]),
      lambda p,n: tc_slide(p,n,[TCARDS[10]]),
      s_appdesign, s_team, s_challenges, s_limits, s_poster, s_conclusion, s_references, s_thanks,
    ]
    for i,b in enumerate(B,1): b(prs,i)
    out=sys.argv[1] if len(sys.argv)>1 else "AEGIS_Presentation.pptx"
    prs.save(out); print("saved",out,"- slides:",len(prs.slides._sldIdLst))

if __name__=="__main__": main()
