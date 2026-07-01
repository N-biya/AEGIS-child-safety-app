from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH

remarks = [
    {
        'form': 11, 'date': '21 April, 2026',
        'points': [
            'Diagrams refined. Sequence diagram updated with BLE flow.',
            'Use-case diagram aligned with defended concept.',
            'Design documentation complete. Ready for implementation phase.',
        ]
    },
    {
        'form': 12, 'date': '28 April, 2026',
        'points': [
            'Firebase connected. Firestore and Auth initialised.',
            'Folder structure and dependencies finalised.',
            'Development roadmap drafted and approved.',
        ]
    },
    {
        'form': 13, 'date': '07 May, 2026',
        'points': [
            'Bottom nav bar implemented. Routing working across all screens.',
            'Page layouts built with consistent styling.',
            'Screens match Velvet Night design. Spacing corrected.',
        ]
    },
    {
        'form': 14, 'date': '13 May, 2026',
        'points': [
            'Login and signup working with Firebase Auth.',
            'Form validation and session persistence in place.',
            'Tested on emulator and physical device. All passing.',
        ]
    },
    {
        'form': 15, 'date': '20 May, 2026',
        'points': [
            'Google Maps integrated. Live location displaying correctly.',
            'Firestore listener active. Coordinates updating in real time.',
            'Custom marker and avatar added to tracking screen.',
        ]
    },
    {
        'form': 16, 'date': '26 May, 2026',
        'points': [
            'Safe zone circle implemented. Draggable and resizable.',
            'Zone data saved to Firestore. Loads correctly on relaunch.',
            'Notification triggers on safe zone breach. Tested successfully.',
        ]
    },
    {
        'form': 17, 'date': '03 June, 2026',
        'points': [
            'ESP32 reading GPS from Neo-6M. NMEA parsing working.',
            'BLE broadcasting device ID and GPS payload. Verified.',
            'Flutter app receiving and parsing BLE data correctly.',
        ]
    },
    {
        'form': 18, 'date': '09 June, 2026',
        'points': [
            'TFLite model running on ESP32. Inference on MPU-6050 confirmed.',
            'Preprocessing pipeline in place. Readings normalised before inference.',
            'Activity labels and fall detection at acceptable accuracy.',
        ]
    },
    {
        'form': 19, 'date': '16 June, 2026',
        'points': [
            'FCM set up. Push notifications reaching parent device.',
            'All three alert types triggering correctly in tests.',
            'Alert history screen complete. Entries showing with timestamps.',
        ]
    },
    {
        'form': 20, 'date': '18 June, 2026',
        'points': [
            'Full BLE pipeline tested. Packet-loss bug found and fixed.',
            'Firestore sync stable under load. Response times acceptable.',
            'All bugs resolved. System ready for final evaluation.',
        ]
    },
]

doc = Document()
section = doc.sections[0]
section.top_margin = Cm(2.0)
section.bottom_margin = Cm(2.0)
section.left_margin = Cm(2.5)
section.right_margin = Cm(2.5)

# Title
t = doc.add_paragraph()
t.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = t.add_run('AEGIS — Child Safety Band')
r.bold = True
r.font.size = Pt(13)

t2 = doc.add_paragraph()
t2.alignment = WD_ALIGN_PARAGRAPH.CENTER
r2 = t2.add_run('Supervisor Remarks — Progress on Previous Items (Forms 11–20)')
r2.bold = True
r2.font.size = Pt(11)

doc.add_paragraph()

for entry in remarks:
    # Meeting heading
    heading = doc.add_paragraph()
    r = heading.add_run(f'Meeting Form {entry["form"]}   |   {entry["date"]}')
    r.bold = True
    r.font.size = Pt(10.5)
    heading.paragraph_format.space_after = Pt(2)

    # Bullet points
    for point in entry['points']:
        p = doc.add_paragraph(style='List Bullet')
        run = p.add_run(point)
        run.font.size = Pt(10)
        p.paragraph_format.space_after = Pt(1)
        p.paragraph_format.space_before = Pt(1)

    doc.add_paragraph().paragraph_format.space_after = Pt(4)

out = 'forms/Progress Notes (Forms 11-20) UPDATED.docx'
try:
    doc.save(out)
    print(f'Saved: {out}')
except PermissionError:
    out = 'forms/Progress Notes (Forms 11-20) v2.docx'
    doc.save(out)
    print(f'Old file was open. Saved as: {out}')
